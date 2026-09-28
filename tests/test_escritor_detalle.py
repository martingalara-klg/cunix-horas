"""El archivo que recibe el partner, releído: formato, orden e integridad.

El contrato es `Horas KLG-Sept2025.xlsx`, el archivo que mandó el partner y
que está en la raíz del repo. Los valores de formato que se verifican acá
(anchos, relleno, formatos de celda) salieron de abrirlo y mirarlo.
"""
from datetime import date, time, timedelta

import openpyxl
import pytest

from cunix_horas.detalle import construir
from cunix_horas.escritor_detalle import (
    ANCHOS,
    ENCABEZADOS,
    RELLENO_ENCABEZADO,
    ErrorIntegridad,
    escribir_detalle,
    marca_de_incompleto,
    nombre_de_archivo,
    nombre_incompleto,
    segundos_escritos,
    verificar_integridad,
)
from cunix_horas.kimai_comun import Registro
from cunix_horas.mapeo import PATRON_ARCHIVO_SALIDA, DestinoProyecto, Mapeo

MAPEO = Mapeo({}, {"CO2510115": DestinoProyecto("ISP de Chile", "SIAC-OIRS")})


def _registro(
    dia=3,
    horas=1.0,
    nombre="Matias Zalazar",
    hora_inicio=time(9, 0),
    descripcion="Ticket R-012528",
):
    return Registro(
        fecha=date(2026, 8, dia),
        horas=horas,
        username="mzalazar",
        cod_proyecto="CO2510115",
        actividad="Desarrollo",
        texto_proyecto="[CO2510115] ISPCH-SopEvo-SIAC | largo",
        hora_inicio=hora_inicio,
        nombre=nombre,
        email="matias.zalazar@cunix.net",
        descripcion=descripcion,
        numero_proyecto="210",
        texto_cliente="[616050001] ISP de Chile",
    )


def _escribir(tmp_path, registros):
    destino = tmp_path / "Horas KLG-Aug2026.xlsx"
    escribir_detalle(construir(registros, MAPEO), destino)
    return openpyxl.load_workbook(destino).active, destino


# --- Formato ----------------------------------------------------------------


def test_las_diez_columnas_van_en_el_orden_del_partner(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro()])
    assert [c.value for c in hoja[1]] == list(ENCABEZADOS)
    assert [c.value for c in hoja[1]] == [
        "Date",
        "Duration",
        "Name",
        "User",
        "E-mail",
        "Customer",
        "Project",
        "Activity",
        "Description",
        "Project number",
    ]


def test_los_encabezados_van_en_negrita_con_su_relleno(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro()])
    for celda in hoja[1]:
        assert celda.font.b is True
        assert celda.fill.patternType == "solid"
        assert celda.fill.fgColor.rgb == RELLENO_ENCABEZADO


def test_los_anchos_son_los_del_archivo_del_partner(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro()])
    anchos = [hoja.column_dimensions[chr(ord("A") + i)].width for i in range(10)]
    assert anchos == list(ANCHOS)
    assert anchos[0] == 9.22 and anchos[9] == 15.0


def test_hay_autofiltro_sobre_el_rango_de_datos(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro(dia=d) for d in (3, 4, 5)])
    assert hoja.auto_filter.ref == "A1:J4"


def test_la_duracion_se_relee_como_timedelta_y_no_como_numero(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro(horas=2.5)])
    assert hoja["B2"].value == timedelta(hours=2, minutes=30)
    assert not isinstance(hoja["B2"].value, (int, float))
    assert hoja["B2"].number_format == "[hh]:mm"


def test_la_fecha_conserva_la_hora_aunque_el_formato_no_la_muestre(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro(dia=5, hora_inicio=time(14, 30))])
    assert hoja["A2"].value.hour == 14
    assert hoja["A2"].value.minute == 30
    assert hoja["A2"].number_format == "yyyy-mm-dd"


def test_una_descripcion_vacia_queda_como_celda_vacia(tmp_path):
    hoja, _ = _escribir(tmp_path, [_registro(descripcion="")])
    assert hoja["I2"].value is None


def test_las_filas_salen_agrupadas_por_dev_y_cronologicas(tmp_path):
    registros = [
        _registro(dia=9, nombre="Matias Zalazar"),
        _registro(dia=2, nombre="Alexis Carnero"),
        _registro(dia=1, nombre="Matias Zalazar"),
    ]
    hoja, _ = _escribir(tmp_path, registros)
    leidas = [
        (hoja.cell(row=f, column=3).value, hoja.cell(row=f, column=1).value.day)
        for f in range(2, 5)
    ]
    assert leidas == [
        ("Alexis Carnero", 2),
        ("Matias Zalazar", 1),
        ("Matias Zalazar", 9),
    ]


# --- Nombre del archivo -----------------------------------------------------


def test_el_nombre_por_defecto_de_agosto_2026():
    assert nombre_de_archivo(PATRON_ARCHIVO_SALIDA, 2026, 8) == (
        "Horas KLG-Aug2026.xlsx"
    )


def test_el_patron_se_puede_cambiar_desde_el_mapeo():
    assert nombre_de_archivo("Horas KLG-Sept{anio}.xlsx", 2025, 9) == (
        "Horas KLG-Sept2025.xlsx"
    )


def test_el_nombre_incompleto_dice_cuantos_faltan():
    assert nombre_incompleto("Horas KLG-Aug2026.xlsx", 2) == (
        "Horas KLG-Aug2026 (INCOMPLETO - FALTAN 2 DESARROLLADORES - NO ENVIAR).xlsx"
    )
    assert "FALTA 1 DESARROLLADOR " in marca_de_incompleto(1)


# --- Integridad -------------------------------------------------------------


def test_la_verificacion_pasa_cuando_el_archivo_tiene_todas_las_horas(tmp_path):
    registros = [_registro(dia=d, horas=1.000000000000008) for d in (3, 4, 5)]
    _, destino = _escribir(tmp_path, registros)
    verificar_integridad(destino, 3 * 3600)
    assert segundos_escritos(destino) == 3 * 3600


def test_la_verificacion_detecta_una_fila_perdida(tmp_path):
    """Con 160 filas en un solo archivo, una fila de menos no se ve a ojo."""
    registros = [_registro(dia=d) for d in (3, 4, 5)]
    _, destino = _escribir(tmp_path, registros)

    libro = openpyxl.load_workbook(destino)
    libro.active.delete_rows(2)
    libro.save(destino)

    with pytest.raises(ErrorIntegridad) as excepcion:
        verificar_integridad(destino, 3 * 3600)
    mensaje = str(excepcion.value)
    assert "2.00 h" in mensaje and "3.00 h" in mensaje
    assert "-1.00 h" in mensaje
    assert "No se generó nada" in mensaje
