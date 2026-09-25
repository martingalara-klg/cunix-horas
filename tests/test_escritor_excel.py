import hashlib

import openpyxl
import pytest
from conftest import FIXTURES

from cunix_horas.agregador import Fila, Reporte
from cunix_horas.escritor_excel import escribir, nombre_de_archivo


def reporte_de_ejemplo(anio=2025, mes=10):
    return Reporte(
        nombre_dev="Franco Dodera",
        nombre_archivo="Dodera",
        anio=anio,
        mes=mes,
        filas=(
            Fila("Club Atlético Talleres", "CRM", "Desarrollo", {4: 2.0, 31: 3.0}),
            Fila("Sistemas - C.UNIX", "Fan Player", "Desarrollo", {31: 4.0}),
        ),
        descartados=(),
    )


@pytest.fixture
def generado(tmp_path):
    destino = tmp_path / "salida.xlsx"
    escribir(reporte_de_ejemplo(), FIXTURES / "plantilla.xlsx", destino)
    return openpyxl.load_workbook(destino).active


def test_nombre_de_archivo():
    assert nombre_de_archivo(reporte_de_ejemplo()) == "Oct Dodera.xlsx"
    assert nombre_de_archivo(reporte_de_ejemplo(mes=1)) == "Jan Dodera.xlsx"


def test_encabezado(generado):
    assert generado["A1"].value == "Franco Dodera"
    assert generado["B1"].value == "Total"
    assert generado["C1"].value == "10/1/2025"
    assert generado["D1"].value == "10/2/2025"
    assert generado["AG1"].value == "10/31/2025"


def test_la_cantidad_de_columnas_sigue_los_dias_del_mes(tmp_path):
    for anio, mes, ultima_columna in [(2025, 2, 30), (2024, 2, 31), (2025, 4, 32)]:
        destino = tmp_path / f"{anio}-{mes}.xlsx"
        escribir(
            Reporte("X", "X", anio, mes, (), ()),
            FIXTURES / "plantilla.xlsx",
            destino,
        )
        hoja = openpyxl.load_workbook(destino).active
        assert hoja.max_column == ultima_columna


def test_jerarquia_de_filas(generado):
    assert generado["A2"].value == "Club Atlético Talleres"
    assert generado["A3"].value == "CRM"
    assert generado["A4"].value == "Desarrollo"
    assert generado["A5"].value == "Sistemas - C.UNIX"
    assert generado["A6"].value == "Fan Player"
    assert generado["A7"].value == "Desarrollo"
    assert generado["A8"].value == "Total"


def test_totales_de_la_columna_b(generado):
    assert generado["B2"].value == 5.0
    assert generado["B3"].value == 5.0
    assert generado["B4"].value == 5.0
    assert generado["B5"].value == 4.0
    assert generado["B8"].value == 9.0


def test_horas_por_dia(generado):
    assert generado["F3"].value == 2.0   # día 4 -> columna 2+4 = F
    assert generado["AG3"].value == 3.0  # día 31 -> columna 33 = AG
    assert generado["AG6"].value == 4.0
    assert generado["D3"].value is None  # día 2, sin horas -> vacío


def test_la_fila_total_pone_cero_en_los_dias_sin_horas(generado):
    assert generado["C8"].value == 0.0
    assert generado["F8"].value == 2.0
    assert generado["AG8"].value == 7.0


def test_las_filas_de_cliente_no_tienen_horas_por_dia(generado):
    assert generado["F2"].value is None


def test_las_filas_de_proyecto_van_en_negrita(generado):
    assert generado["A3"].font.b is True
    assert generado["A6"].font.b is True
    assert generado["A2"].font.b is not True
    assert generado["A4"].font.b is not True


def test_merge_en_las_filas_de_cliente(generado):
    rangos = {str(r) for r in generado.merged_cells.ranges}
    assert "C2:AG2" in rangos
    assert "C5:AG5" in rangos


def test_anchos_de_columna_copiados_de_la_plantilla(generado):
    assert generado.column_dimensions["A"].width == pytest.approx(34.14, abs=0.01)
    assert generado.column_dimensions["B"].width == pytest.approx(9.29, abs=0.01)


def test_el_nombre_de_la_hoja_es_el_de_la_plantilla(generado):
    assert generado.title == "Worksheet"


def test_las_horas_se_redondean_a_dos_decimales(tmp_path):
    destino = tmp_path / "r.xlsx"
    escribir(
        Reporte("X", "X", 2025, 10, (Fila("C", "P", "A", {1: 1 / 3}),), ()),
        FIXTURES / "plantilla.xlsx",
        destino,
    )
    hoja = openpyxl.load_workbook(destino).active
    assert hoja["C4"].value == 0.33


# --- Regresión: la negrita se aplica por celda, no por fila entera ---
# (defecto Critical: el diseño original tomaba UN estilo por tipo de fila y lo
# aplicaba a toda la fila; en la plantilla real la negrita de la fila de
# proyecto sólo va en el nombre y en los días CON horas, no en el total ni en
# los días vacíos.)


def test_negrita_por_celda_en_fila_de_proyecto(generado):
    # Fila 3 = proyecto "CRM": A3 nombre (negrita), B3 total (no), F3 día
    # con horas -> día 4 (negrita), D3 día sin horas -> día 2 (no).
    assert generado["A3"].font.b is True
    assert generado["B3"].font.b is not True
    assert generado["F3"].font.b is True
    assert generado["D3"].font.b is not True


def test_filas_de_cliente_actividad_y_total_no_tienen_ninguna_celda_en_negrita(generado):
    for nro_fila in (2, 4, 5, 7, 8):  # cliente, actividad, cliente, actividad, total
        for columna in range(1, generado.max_column + 1):
            celda = generado.cell(row=nro_fila, column=columna)
            assert celda.font.b is not True, f"{celda.coordinate} no debería ir en negrita"


def _reporte_de_la_plantilla_fixture():
    """Reconstruye el Reporte cuyos datos de horas produjeron
    tests/fixtures/plantilla.xlsx (Sistemas - C.UNIX / Fan Player / Desarrollo
    y Club Atlético Talleres / CRM / Desarrollo), en el mismo orden en que
    aparecen las filas de cliente en el archivo (Sistemas primero)."""
    return Reporte(
        nombre_dev="Franco Dodera",
        nombre_archivo="Dodera",
        anio=2025,
        mes=10,
        filas=(
            Fila(
                "Sistemas - C.UNIX",
                "Fan Player",
                "Desarrollo",
                {8: 6.0, 9: 4.0, 15: 4.0, 16: 4.5, 21: 4.0, 22: 4.0, 23: 6.0,
                 24: 7.0, 29: 5.0, 30: 7.0, 31: 4.0},
            ),
            Fila(
                "Club Atlético Talleres",
                "CRM",
                "Desarrollo",
                {2: 2.0, 3: 3.0, 7: 5.0, 8: 6.0, 9: 3.0, 13: 7.0, 14: 4.0,
                 15: 5.0, 16: 4.5, 21: 3.0, 22: 4.0, 23: 3.0, 24: 3.0,
                 27: 6.0, 28: 5.0, 29: 3.0, 30: 3.0, 31: 3.0},
            ),
        ),
        descartados=(),
    )


def test_negrita_coincide_celda_por_celda_con_la_plantilla(tmp_path):
    # Compara SOLO el patrón de negrita, nunca los valores: la fila 'Total'
    # de la fixture está recortada y sus valores no cierran con las filas de
    # arriba.
    destino = tmp_path / "comparado.xlsx"
    escribir(_reporte_de_la_plantilla_fixture(), FIXTURES / "plantilla.xlsx", destino)
    generado = openpyxl.load_workbook(destino).active
    esperado = openpyxl.load_workbook(FIXTURES / "plantilla.xlsx").active

    assert generado.max_row == esperado.max_row
    assert generado.max_column == esperado.max_column
    for nro_fila in range(1, esperado.max_row + 1):
        for columna in range(1, esperado.max_column + 1):
            real = generado.cell(row=nro_fila, column=columna)
            modelo = esperado.cell(row=nro_fila, column=columna)
            assert bool(real.font.b) == bool(modelo.font.b), (
                f"negrita distinta en {real.coordinate}: "
                f"generado={real.font.b!r} plantilla={modelo.font.b!r}"
            )


def test_cliente_con_dos_proyectos_y_proyecto_con_dos_actividades(tmp_path):
    reporte = Reporte(
        nombre_dev="X",
        nombre_archivo="X",
        anio=2025,
        mes=10,
        filas=(
            Fila("Cliente", "ProyectoA", "Act1", {1: 1.0, 2: 2.0}),
            Fila("Cliente", "ProyectoA", "Act2", {1: 3.0}),
            Fila("Cliente", "ProyectoB", "Act1", {5: 4.0}),
        ),
        descartados=(),
    )
    destino = tmp_path / "multi.xlsx"
    escribir(reporte, FIXTURES / "plantilla.xlsx", destino)
    hoja = openpyxl.load_workbook(destino).active

    # Filas: 1 encabezado, 2 cliente, 3 ProyectoA, 4 Act1, 5 Act2,
    # 6 ProyectoB, 7 Act1, 8 Total.
    assert hoja["A2"].value == "Cliente"
    assert hoja["B2"].value == pytest.approx(10.0)  # total cliente = 6.0 + 4.0

    assert hoja["A3"].value == "ProyectoA"
    assert hoja["B3"].value == pytest.approx(6.0)  # (1+2) + 3
    assert hoja["C3"].value == pytest.approx(4.0)  # día 1: Act1 1.0 + Act2 3.0
    assert hoja["D3"].value == pytest.approx(2.0)  # día 2: sólo Act1

    assert hoja["A4"].value == "Act1"
    assert hoja["A5"].value == "Act2"

    assert hoja["A6"].value == "ProyectoB"
    assert hoja["B6"].value == pytest.approx(4.0)

    assert hoja["A7"].value == "Act1"
    assert hoja["A8"].value == "Total"
    assert hoja["B8"].value == pytest.approx(10.0)


def test_reporte_sin_filas_genera_excel_valido(tmp_path):
    destino = tmp_path / "vacio.xlsx"
    resultado = escribir(
        Reporte("X", "X", 2025, 10, (), ()), FIXTURES / "plantilla.xlsx", destino
    )
    assert resultado == destino
    hoja = openpyxl.load_workbook(destino).active
    assert hoja["A1"].value == "X"
    assert hoja["A2"].value == "Total"
    assert hoja["B2"].value == 0.0


def test_la_plantilla_no_se_modifica(tmp_path):
    plantilla = FIXTURES / "plantilla.xlsx"
    hash_antes = hashlib.sha256(plantilla.read_bytes()).hexdigest()
    escribir(reporte_de_ejemplo(), plantilla, tmp_path / "salida.xlsx")
    hash_despues = hashlib.sha256(plantilla.read_bytes()).hexdigest()
    assert hash_antes == hash_despues


def test_falla_ruidosamente_si_la_plantilla_no_tiene_fila_total(tmp_path):
    libro = openpyxl.load_workbook(FIXTURES / "plantilla.xlsx")
    libro.active["A8"] = "No es total"
    plantilla_rota = tmp_path / "rota.xlsx"
    libro.save(plantilla_rota)

    with pytest.raises(ValueError, match="Total"):
        escribir(reporte_de_ejemplo(), plantilla_rota, tmp_path / "salida.xlsx")
