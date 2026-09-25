import shutil

import openpyxl
import yaml
from conftest import FIXTURES

import cunix_horas.cli as cli
from cunix_horas.cli import procesar_mes
from cunix_horas.escritor_excel import escribir as escribir_real


def preparar(tmp_path, nombres=("kimai-mzalazar.xlsx",)):
    (tmp_path / "config").mkdir()
    (tmp_path / "templates").mkdir()
    entrada = tmp_path / "input" / "2026-08"
    entrada.mkdir(parents=True)
    shutil.copy(FIXTURES / "mapeo-test.yaml", tmp_path / "config" / "mapeo.yaml")
    shutil.copy(FIXTURES / "plantilla.xlsx", tmp_path / "templates" / "plantilla.xlsx")
    for nombre in nombres:
        shutil.copy(FIXTURES / "kimai-mzalazar.xlsx", entrada / nombre)
    return tmp_path


def test_genera_el_excel_del_mes(tmp_path):
    raiz = preparar(tmp_path)
    assert procesar_mes("2026-08", raiz) == 0
    salida = raiz / "output" / "2026-08" / "Aug Zalazar.xlsx"
    assert salida.is_file()
    hoja = openpyxl.load_workbook(salida).active
    assert hoja["A1"].value == "Matias Zalazar"
    assert hoja["B1"].value == "Total"
    assert hoja.cell(row=hoja.max_row, column=1).value == "Total"


def test_los_totales_del_excel_cierran_en_76_5(tmp_path):
    raiz = preparar(tmp_path)
    procesar_mes("2026-08", raiz)
    hoja = openpyxl.load_workbook(
        raiz / "output" / "2026-08" / "Aug Zalazar.xlsx"
    ).active
    assert hoja.cell(row=hoja.max_row, column=2).value == 76.5


def test_escribe_el_archivo_de_validacion(tmp_path):
    raiz = preparar(tmp_path)
    procesar_mes("2026-08", raiz)
    validacion = raiz / "output" / "2026-08" / "_validacion.txt"
    assert validacion.is_file()
    assert "Día hábil sin carga" in validacion.read_text(encoding="utf-8")


def test_un_archivo_roto_no_frena_a_los_demas(tmp_path, capsys):
    raiz = preparar(tmp_path)
    (raiz / "input" / "2026-08" / "roto.xlsx").write_text("basura", encoding="utf-8")
    codigo = procesar_mes("2026-08", raiz)
    assert (raiz / "output" / "2026-08" / "Aug Zalazar.xlsx").is_file()
    assert codigo == 1
    assert "roto.xlsx" in capsys.readouterr().out


def test_falla_si_no_existe_la_carpeta_del_mes(tmp_path, capsys):
    raiz = preparar(tmp_path)
    assert procesar_mes("2026-09", raiz) == 1
    assert "input/2026-09" in capsys.readouterr().out


def test_falla_con_un_mes_mal_escrito(tmp_path, capsys):
    raiz = preparar(tmp_path)
    assert procesar_mes("agosto", raiz) == 1
    assert "AAAA-MM" in capsys.readouterr().out


def test_un_proyecto_sin_mapear_frena_solo_a_ese_dev(tmp_path, capsys):
    raiz = preparar(tmp_path)
    (raiz / "config" / "mapeo.yaml").write_text(
        "personas:\n"
        "  mzalazar:\n"
        '    nombre: "Matias Zalazar"\n'
        '    archivo: "Zalazar"\n'
        "proyectos:\n"
        "  CO2610170:\n"
        '    cliente: "Aduanas"\n'
        '    proyecto: "Subastas"\n',
        encoding="utf-8",
    )
    assert procesar_mes("2026-08", raiz) == 1
    salida = capsys.readouterr().out
    assert "CO2510115" in salida
    assert "config/mapeo.yaml" in salida
    assert not (raiz / "output" / "2026-08" / "Aug Zalazar.xlsx").exists()


def test_el_bloque_yaml_sugerido_en_consola_es_pegable(tmp_path, capsys):
    """Regresión: cli.py no debe romper el sangrado que arma mapeo.py.

    Tarea 3 corrigió el sangrado del YAML sugerido para que fuera pegable
    tal cual. cli.py lo rompía de nuevo indentando cada línea 4 espacios
    más. Este test cubre el camino completo: lo que realmente se imprime
    en consola, no el mensaje crudo de la excepción.
    """
    raiz = preparar(tmp_path)
    (raiz / "config" / "mapeo.yaml").write_text(
        "personas:\n"
        "  mzalazar:\n"
        '    nombre: "Matias Zalazar"\n'
        '    archivo: "Zalazar"\n'
        "proyectos:\n"
        "  CO2610170:\n"
        '    cliente: "Aduanas"\n'
        '    proyecto: "Subastas"\n',
        encoding="utf-8",
    )
    assert procesar_mes("2026-08", raiz) == 1
    salida = capsys.readouterr().out

    lineas = salida.splitlines()
    inicio_bloque = next(i for i, linea in enumerate(lineas) if "CO2510115:" in linea)
    lineas_del_bloque = []
    for linea in lineas[inicio_bloque:]:
        if linea and not linea.startswith(" "):
            break
        lineas_del_bloque.append(linea)
    bloque_yaml = "\n".join(lineas_del_bloque)

    ruta_mapeo = raiz / "config" / "mapeo.yaml"
    ruta_mapeo.write_text(
        ruta_mapeo.read_text(encoding="utf-8") + bloque_yaml + "\n",
        encoding="utf-8",
    )

    datos = yaml.safe_load(ruta_mapeo.read_text(encoding="utf-8"))
    assert "CO2510115" in datos["proyectos"]
    assert isinstance(datos["proyectos"]["CO2510115"], dict)
    assert "cliente" in datos["proyectos"]["CO2510115"]
    assert "proyecto" in datos["proyectos"]["CO2510115"]


def test_dos_archivos_que_resuelven_al_mismo_nombre_no_se_pisan(tmp_path, capsys):
    """Dos exports del mismo dev (o un re-export) no deben pisarse en silencio."""
    raiz = preparar(
        tmp_path, nombres=("kimai-mzalazar-v1.xlsx", "kimai-mzalazar-v2.xlsx")
    )
    codigo = procesar_mes("2026-08", raiz)
    assert codigo == 1
    salida = capsys.readouterr().out
    assert "kimai-mzalazar-v1.xlsx" in salida
    assert "kimai-mzalazar-v2.xlsx" in salida
    assert "Aug Zalazar.xlsx" in salida
    # El primero de los dos igual se genera; no se pierde ese reporte.
    assert (raiz / "output" / "2026-08" / "Aug Zalazar.xlsx").is_file()


def test_un_excel_abierto_no_frena_a_los_demas(tmp_path, capsys, monkeypatch):
    """Si escribir() falla por PermissionError (Excel abierto), no debe caer todo."""
    raiz = preparar(tmp_path, nombres=("kimai-a.xlsx", "kimai-b.xlsx"))

    llamadas = {"n": 0}

    def escribir_simulado(reporte, plantilla, destino):
        llamadas["n"] += 1
        if llamadas["n"] == 1:
            raise PermissionError(13, "Acceso denegado")
        return escribir_real(reporte, plantilla, destino)

    monkeypatch.setattr(cli, "escribir", escribir_simulado)

    codigo = procesar_mes("2026-08", raiz)
    assert codigo == 1
    salida = capsys.readouterr().out
    assert "kimai-a.xlsx" in salida
    assert "abierto" in salida.lower()
    # El segundo archivo (mismo nombre de salida) sí se generó.
    assert (raiz / "output" / "2026-08" / "Aug Zalazar.xlsx").is_file()
