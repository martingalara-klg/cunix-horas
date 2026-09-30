# -*- coding: utf-8 -*-
"""Genera el respaldo mensual de facturación de AGOSTO 2026, y sólo ése.

Esto NO es parte de la herramienta. `cunix_horas/` produce el Excel que se le
manda al partner todos los meses; este script produce, una sola vez, el Word de
respaldo de agosto de 2026 a partir de la plantilla que usa el dueño. Está
versionado para poder repetirlo tal cual o auditar de dónde salió cada número,
no para correrlo todos los meses.

Por qué la corrección de Gabriel Denis se aplica acá y no en la herramienta
-------------------------------------------------------------------------
Gabriel no carga horas en Kimai: durante agosto de 2026 sus horas se cargaron
en la cuenta de Alexis Carnero. La cuenta de Alexis quedó entonces con 150 h
que son de dos personas. Gabriel lleva su propia planilla (91 h), así que la
separación se hace restando, día por día, lo de Gabriel de lo de Alexis.

Eso es un hecho de un mes puntual, no una regla del negocio:

- no hay nada en Kimai que diga qué parte de un día es de quién; el único dato
  que lo permite es una planilla externa que existe para agosto de 2026;
- meterlo en `cunix_horas/` obligaría a la herramienta a conocer nombres
  propios y a depender de un archivo que el mes que viene puede no existir;
- si el arreglo se vuelve permanente, lo correcto es que Gabriel tenga usuario
  en Kimai, no que el programa reparta horas ajenas.

Por eso la resta vive acá, a la vista, con su verificación propia: si algún día
quedara negativo, si Gabriel tuviera un día que Alexis no tiene, o si los
totales no cerraran, el script aborta y no escribe nada.

Uso:
    python scripts/respaldo_agosto_2026.py
"""
from __future__ import annotations

import copy
import re
import shutil
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import docx  # noqa: E402  (después de arreglar sys.path)
import openpyxl  # noqa: E402

from cunix_horas.completado import completar_desde_mapeo  # noqa: E402
from cunix_horas.detalle import segundos_de  # noqa: E402
from cunix_horas.lector_kimai import leer  # noqa: E402
from cunix_horas.mapeo import (  # noqa: E402
    Mapeo,
    derivar_cliente,
    derivar_proyecto,
)

# ---------------------------------------------------------------------------
# Lo que el dueño va a querer tocar
# ---------------------------------------------------------------------------

# Va en "Perfil Técnico" (resumen) y en "Perfil / Cargo" (anexo), en TODAS las
# filas. Si mañana hay que distinguir seniorities, se cambia acá.
PERFIL_TECNICO = "Desarrollador"

PERIODO = "Agosto 2026"

# ---------------------------------------------------------------------------
# Constantes del mes
# ---------------------------------------------------------------------------

ANIO = 2026
MES = 8

SEGUNDOS_POR_HORA = 3600

PLANTILLA = RAIZ / "Plantilla Respaldo Mensual Facturacion - Editable.docx"
CARPETA_KIMAI = RAIZ / "input" / "2026-08"
MAPEO = RAIZ / "config" / "mapeo.yaml"
PLANILLA_GABRIEL = RAIZ / "horas_agosto_2026 (2).xlsx"
HOJA_GABRIEL = "Agosto 2026"
PRIMERA_FILA_GABRIEL = 5
ULTIMA_FILA_GABRIEL = 20  # De la 22 en adelante son subtotales y total: no entran.

DESTINO = RAIZ / "output" / "2026-08" / "Respaldo Mensual Facturacion - Agosto 2026.docx"

# La persona en cuya cuenta de Kimai quedaron cargadas las horas ajenas, y la
# persona dueña de esas horas.
NOMBRE_ALEXIS = "Alexis Carnero"
NOMBRE_GABRIEL = "Gabriel Denis"

# El proyecto en el que Gabriel trabajó todo el mes. Es el mismo código de
# Kimai con el que están cargadas las horas de Alexis, así que se resuelve por
# el mapeo y no se escribe el nombre a mano.
CODIGO_PROYECTO_GABRIEL = "PR2510126"

TOTAL_ESPERADO_HORAS = 306.0
TOTAL_GABRIEL_HORAS = 91.0
TOTAL_ALEXIS_NETO_HORAS = 59.0
FILAS_ESPERADAS_GABRIEL = 16

# Los dos días que fueron enteramente de Gabriel: Alexis no debe tener fila.
DIAS_SIN_ALEXIS = ("02/08/2026", "20/08/2026")

OBS_RESUMEN_ALEXIS = (
    "Horas netas de Alexis Carnero: durante agosto las horas de Gabriel Denis "
    "quedaron registradas en esta cuenta de Kimai y se separaron día por día."
)
OBS_RESUMEN_GABRIEL = (
    "Horas tomadas del registro propio de Gabriel Denis: estaban cargadas en la "
    "cuenta de Alexis Carnero en Kimai y se separaron día por día."
)
OBS_ANEXO_ALEXIS = "Horas netas, ya descontadas las de Gabriel Denis."

# Índices de las tablas de la plantilla.
TABLA_GENERALES = 0
TABLA_RESUMEN = 1
TABLA_ANEXO = 3
# Firmas y checklist: el script no las toca y la verificación lo comprueba.
TABLAS_INTACTAS = (2, 4)

_DIA_EN_TEXTO = re.compile(r"(\d{1,2})")


class ErrorDeRespaldo(Exception):
    """Algo no cierra: se aborta sin escribir nada."""


# ---------------------------------------------------------------------------
# Datos
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Linea:
    """Un registro de horas ya listo para el documento.

    Las horas viajan en segundos enteros: las duraciones de Kimai llegan como
    float (1 h vale 1.000000000000008) y sumarlas así arrastra el error a un
    documento de facturación.
    """

    persona: str
    fecha: date
    cliente: str
    proyecto: str
    descripcion: str
    segundos: int
    observaciones: str = ""

    @property
    def proyecto_para_mostrar(self) -> str:
        return f"{self.cliente} / {self.proyecto}"


def _horas(segundos: int) -> float:
    return segundos / SEGUNDOS_POR_HORA


def _texto_horas(segundos: int) -> str:
    """Las horas como texto, sin perder nada por redondeo.

    La mayoría de los registros son horas o medias horas y salen con un
    decimal ('8.0', '5.5'), pero hay cuartos de hora: con un solo decimal,
    0.25 se escribiría '0.2' y el documento dejaría de sumar. Por eso se usan
    los decimales que hagan falta, y se comprueba que el texto escrito
    represente exactamente los segundos del registro.
    """
    for decimales in (1, 2, 4):
        texto = f"{_horas(segundos):.{decimales}f}"
        if round(float(texto) * SEGUNDOS_POR_HORA) == segundos:
            return texto
    raise ErrorDeRespaldo(
        f"No se puede escribir {segundos} segundos como horas sin perder "
        f"precisión: el documento dejaría de sumar."
    )


def _nombre_de(registro) -> str:
    return registro.nombre.strip() or registro.username.strip()


def _cliente_y_proyecto(mapeo: Mapeo, registro) -> tuple[str, str]:
    """Mismo criterio que usa `cunix_horas.detalle` para el Excel del partner."""
    destino = mapeo.proyecto_opcional(registro.cod_proyecto)
    if destino is not None:
        return destino.cliente, destino.proyecto
    cliente = derivar_cliente(registro.texto_cliente)
    proyecto = derivar_proyecto(registro.texto_proyecto) or registro.cod_proyecto
    return cliente, proyecto


def leer_kimai(mapeo: Mapeo) -> list[Linea]:
    """Los cinco exports de input/2026-08/, con el lector del propio proyecto."""
    lineas: list[Linea] = []
    for archivo in sorted(CARPETA_KIMAI.iterdir()):
        if archivo.suffix.lower() not in (".xlsx", ".csv"):
            continue
        registros = completar_desde_mapeo(leer(archivo), mapeo, archivo.name)
        for registro in registros:
            if (registro.fecha.year, registro.fecha.month) != (ANIO, MES):
                raise ErrorDeRespaldo(
                    f"{archivo.name} trae un registro de {registro.fecha}, que no "
                    f"es de agosto de 2026."
                )
            cliente, proyecto = _cliente_y_proyecto(mapeo, registro)
            lineas.append(
                Linea(
                    persona=_nombre_de(registro),
                    fecha=registro.fecha,
                    cliente=cliente,
                    proyecto=proyecto,
                    descripcion=registro.descripcion.strip(),
                    segundos=segundos_de(registro.horas),
                )
            )
    if not lineas:
        raise ErrorDeRespaldo(f"No se leyó ningún registro de {CARPETA_KIMAI}.")
    return lineas


def leer_gabriel(mapeo: Mapeo) -> list[Linea]:
    """La planilla propia de Gabriel: filas 5 a 20 de la hoja 'Agosto 2026'.

    La fecha viene como texto en español y sin año ('Mié 05 Ago'): el número
    del día es lo único aprovechable, y el mes y el año son los del período.
    """
    destino = mapeo.proyecto_opcional(CODIGO_PROYECTO_GABRIEL)
    if destino is None:
        raise ErrorDeRespaldo(
            f"El proyecto {CODIGO_PROYECTO_GABRIEL} no está en {MAPEO}: sin él no "
            f"se puede nombrar el proyecto de Gabriel Denis."
        )

    libro = openpyxl.load_workbook(PLANILLA_GABRIEL, data_only=True)
    if HOJA_GABRIEL not in libro.sheetnames:
        raise ErrorDeRespaldo(
            f"{PLANILLA_GABRIEL.name} no tiene la hoja '{HOJA_GABRIEL}'."
        )
    hoja = libro[HOJA_GABRIEL]

    lineas: list[Linea] = []
    for numero in range(PRIMERA_FILA_GABRIEL, ULTIMA_FILA_GABRIEL + 1):
        texto_fecha = hoja.cell(row=numero, column=1).value
        horas = hoja.cell(row=numero, column=2).value
        descripcion = hoja.cell(row=numero, column=3).value

        if texto_fecha is None and horas is None:
            continue
        if texto_fecha is None or horas is None:
            raise ErrorDeRespaldo(
                f"Fila {numero} de {PLANILLA_GABRIEL.name}: tiene fecha o horas, "
                f"pero no las dos (fecha={texto_fecha!r}, horas={horas!r})."
            )

        coincidencia = _DIA_EN_TEXTO.search(str(texto_fecha))
        if coincidencia is None:
            raise ErrorDeRespaldo(
                f"Fila {numero} de {PLANILLA_GABRIEL.name}: no se encontró el día "
                f"en la fecha {texto_fecha!r}."
            )
        dia = int(coincidencia.group(1))
        try:
            fecha = date(ANIO, MES, dia)
        except ValueError as error:
            raise ErrorDeRespaldo(
                f"Fila {numero} de {PLANILLA_GABRIEL.name}: día {dia} inválido "
                f"para agosto de 2026."
            ) from error

        segundos = segundos_de(float(horas))
        if segundos <= 0:
            raise ErrorDeRespaldo(
                f"Fila {numero} de {PLANILLA_GABRIEL.name}: horas {horas!r} no "
                f"positivas."
            )

        lineas.append(
            Linea(
                persona=NOMBRE_GABRIEL,
                fecha=fecha,
                cliente=destino.cliente,
                proyecto=destino.proyecto,
                descripcion=str(descripcion or "").strip(),
                segundos=segundos,
            )
        )

    if not lineas:
        raise ErrorDeRespaldo(
            f"No se leyó ninguna fila de horas de {PLANILLA_GABRIEL.name}."
        )
    return lineas


# ---------------------------------------------------------------------------
# La corrección
# ---------------------------------------------------------------------------


def separar_horas_de_gabriel(
    lineas_kimai: list[Linea], lineas_gabriel: list[Linea]
) -> list[Linea]:
    """Resta, día por día, las horas de Gabriel de la cuenta de Alexis.

    Aborta si algo no cierra. No se da por supuesto nada de lo que dijo el
    dueño: esto es la base de una factura.
    """
    de_alexis = [linea for linea in lineas_kimai if linea.persona == NOMBRE_ALEXIS]
    resto = [linea for linea in lineas_kimai if linea.persona != NOMBRE_ALEXIS]
    if not de_alexis:
        raise ErrorDeRespaldo(
            f"No hay ningún registro de {NOMBRE_ALEXIS} en los exports de Kimai: "
            f"no hay de dónde descontar las horas de {NOMBRE_GABRIEL}."
        )

    proyectos = {(linea.cliente, linea.proyecto) for linea in de_alexis}
    if len(proyectos) > 1:
        raise ErrorDeRespaldo(
            f"{NOMBRE_ALEXIS} tiene horas en más de un proyecto "
            f"({sorted(proyectos)}): no se puede saber de cuál descontar las "
            f"horas de {NOMBRE_GABRIEL} sin inventar el reparto."
        )
    cliente_alexis, proyecto_alexis = proyectos.pop()

    por_dia_alexis: dict[date, int] = defaultdict(int)
    for linea in de_alexis:
        por_dia_alexis[linea.fecha] += linea.segundos

    por_dia_gabriel: dict[date, int] = defaultdict(int)
    for linea in lineas_gabriel:
        por_dia_gabriel[linea.fecha] += linea.segundos

    faltantes = sorted(set(por_dia_gabriel) - set(por_dia_alexis))
    if faltantes:
        dias = ", ".join(f.strftime("%d/%m/%Y") for f in faltantes)
        raise ErrorDeRespaldo(
            f"{NOMBRE_GABRIEL} declara horas en días en los que {NOMBRE_ALEXIS} "
            f"no tiene ninguna cargada: {dias}. La resta no se puede hacer."
        )

    negativos: list[str] = []
    netas: list[Linea] = []
    for fecha in sorted(por_dia_alexis):
        restantes = por_dia_alexis[fecha] - por_dia_gabriel.get(fecha, 0)
        if restantes < 0:
            negativos.append(
                f"{fecha.strftime('%d/%m/%Y')} "
                f"({_texto_horas(por_dia_alexis[fecha])} h de {NOMBRE_ALEXIS} "
                f"contra {_texto_horas(por_dia_gabriel[fecha])} h de "
                f"{NOMBRE_GABRIEL})"
            )
            continue
        if restantes == 0:
            # Ese día fue enteramente de Gabriel: Alexis no lleva fila.
            continue
        netas.append(
            Linea(
                persona=NOMBRE_ALEXIS,
                fecha=fecha,
                cliente=cliente_alexis,
                proyecto=proyecto_alexis,
                descripcion="",
                segundos=restantes,
                observaciones=OBS_ANEXO_ALEXIS,
            )
        )

    if negativos:
        raise ErrorDeRespaldo(
            f"La resta deja días en negativo: {'; '.join(negativos)}. "
            f"Revisá la planilla de {NOMBRE_GABRIEL} contra Kimai antes de "
            f"generar el respaldo."
        )

    total_alexis_bruto = sum(por_dia_alexis.values())
    total_gabriel = sum(por_dia_gabriel.values())
    total_neto = sum(linea.segundos for linea in netas)
    if total_neto + total_gabriel != total_alexis_bruto:
        raise ErrorDeRespaldo(
            f"Los totales no cierran: {NOMBRE_ALEXIS} bruto "
            f"{_texto_horas(total_alexis_bruto)} h, {NOMBRE_GABRIEL} "
            f"{_texto_horas(total_gabriel)} h, neto de {NOMBRE_ALEXIS} "
            f"{_texto_horas(total_neto)} h."
        )
    if _horas(total_gabriel) != TOTAL_GABRIEL_HORAS:
        raise ErrorDeRespaldo(
            f"{NOMBRE_GABRIEL} suma {_texto_horas(total_gabriel)} h y se "
            f"esperaban {TOTAL_GABRIEL_HORAS}."
        )
    if _horas(total_neto) != TOTAL_ALEXIS_NETO_HORAS:
        raise ErrorDeRespaldo(
            f"{NOMBRE_ALEXIS} queda con {_texto_horas(total_neto)} h netas y se "
            f"esperaban {TOTAL_ALEXIS_NETO_HORAS}."
        )

    return resto + netas + lineas_gabriel


# ---------------------------------------------------------------------------
# Armado de las tablas
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class FilaResumen:
    persona: str
    proyecto_para_mostrar: str
    segundos: int
    observaciones: str


def armar_resumen(lineas: list[Linea]) -> list[FilaResumen]:
    """Una fila por persona y proyecto, alfabético por persona."""
    acumulado: dict[tuple[str, str], int] = defaultdict(int)
    for linea in lineas:
        acumulado[(linea.persona, linea.proyecto_para_mostrar)] += linea.segundos

    observaciones = {
        NOMBRE_ALEXIS: OBS_RESUMEN_ALEXIS,
        NOMBRE_GABRIEL: OBS_RESUMEN_GABRIEL,
    }

    filas = [
        FilaResumen(
            persona=persona,
            proyecto_para_mostrar=proyecto,
            segundos=segundos,
            observaciones=observaciones.get(persona, ""),
        )
        for (persona, proyecto), segundos in acumulado.items()
    ]
    # Alfabético por persona y, dentro de cada una, de más horas a menos.
    filas.sort(
        key=lambda f: (f.persona.casefold(), -f.segundos, f.proyecto_para_mostrar)
    )
    return filas


def armar_anexo(lineas: list[Linea]) -> list[Linea]:
    """Por persona (alfabético) y, dentro de cada una, por fecha."""
    return sorted(
        lineas,
        key=lambda linea: (
            linea.persona.casefold(),
            linea.fecha,
            linea.proyecto_para_mostrar,
            linea.descripcion,
        ),
    )


# ---------------------------------------------------------------------------
# Escritura del .docx
# ---------------------------------------------------------------------------


def _escribir_celda(celda, texto: str) -> None:
    """Pone `texto` en la celda conservando el formato de la plantilla.

    Las celdas vacías de la plantilla traen un `w:r` con su `w:rPr` (Arial,
    tamaño 8.5): se escribe sobre el último run y se vacían los demás, en vez
    de reemplazar el párrafo, para no perder fuente ni alineación.
    """
    parrafo = celda.paragraphs[0]
    for sobrante in celda.paragraphs[1:]:
        sobrante._element.getparent().remove(sobrante._element)
    if not parrafo.runs:
        parrafo.add_run("")
    for run in parrafo.runs[:-1]:
        run.text = ""
    parrafo.runs[-1].text = texto


def _agregar_filas(tabla, cuantas: int, indice_modelo: int, indice_antes_de: int) -> None:
    """Clona `cuantas` veces una fila existente y las inserta antes del total.

    Se clona el XML de una fila de la plantilla para que las nuevas hereden
    bordes, alto, anchos y fuentes, en vez de salir crudas y sin estilo.
    """
    modelo = tabla.rows[indice_modelo]._tr
    ancla = tabla.rows[indice_antes_de]._tr
    for _ in range(cuantas):
        ancla.addprevious(copy.deepcopy(modelo))


def _preparar_filas(tabla, necesarias: int) -> int:
    """Deja la tabla con al menos `necesarias` filas de datos. Devuelve el total."""
    fila_total = len(tabla.rows) - 1
    disponibles = fila_total - 1  # entre el encabezado y la fila de total
    if necesarias > disponibles:
        _agregar_filas(
            tabla,
            necesarias - disponibles,
            indice_modelo=1,
            indice_antes_de=fila_total,
        )
    return len(tabla.rows) - 1


def completar_documento(
    documento, filas_resumen: list[FilaResumen], filas_anexo: list[Linea]
) -> None:
    generales = documento.tables[TABLA_GENERALES]
    _escribir_celda(generales.rows[0].cells[1], PERIODO)

    # --- Resumen de horas y cobro -----------------------------------------
    resumen = documento.tables[TABLA_RESUMEN]
    fila_total = _preparar_filas(resumen, len(filas_resumen))
    for numero, fila in enumerate(filas_resumen, start=1):
        celdas = resumen.rows[numero].cells
        _escribir_celda(celdas[0], PERFIL_TECNICO)
        _escribir_celda(celdas[1], _texto_horas(fila.segundos))
        _escribir_celda(celdas[2], "")  # Tarifa por Hora: la pone el dueño.
        _escribir_celda(celdas[3], "")  # Subtotal: sale de la tarifa.
        _escribir_celda(celdas[4], fila.proyecto_para_mostrar)
        _escribir_celda(celdas[5], fila.persona)
        _escribir_celda(celdas[6], fila.observaciones)
    total_resumen = sum(fila.segundos for fila in filas_resumen)
    _escribir_celda(resumen.rows[fila_total].cells[1], _texto_horas(total_resumen))

    # --- Anexo detallado ---------------------------------------------------
    anexo = documento.tables[TABLA_ANEXO]
    fila_total_anexo = _preparar_filas(anexo, len(filas_anexo))
    for numero, linea in enumerate(filas_anexo, start=1):
        celdas = anexo.rows[numero].cells
        _escribir_celda(celdas[0], linea.fecha.strftime("%d/%m/%Y"))
        _escribir_celda(celdas[1], linea.proyecto_para_mostrar)
        _escribir_celda(celdas[2], linea.persona)
        _escribir_celda(celdas[3], PERFIL_TECNICO)
        _escribir_celda(celdas[4], "")  # Ticket / Tarea: no se deduce del texto.
        _escribir_celda(celdas[5], linea.descripcion)
        _escribir_celda(celdas[6], _texto_horas(linea.segundos))
        _escribir_celda(celdas[7], linea.observaciones)
    total_anexo = sum(linea.segundos for linea in filas_anexo)
    _escribir_celda(anexo.rows[fila_total_anexo].cells[6], _texto_horas(total_anexo))


# ---------------------------------------------------------------------------
# Verificación sobre el archivo ya escrito
# ---------------------------------------------------------------------------


def _filas_con_datos(tabla) -> list[list[str]]:
    """El texto de las filas de datos: sin encabezado, sin total y sin vacías."""
    filas = [
        [celda.text.strip() for celda in fila.cells] for fila in tabla.rows[1:-1]
    ]
    return [fila for fila in filas if any(fila)]


def verificar_documento(ruta: Path) -> str:
    """Reabre el .docx generado y comprueba los números leyéndolo."""
    documento = docx.Document(str(ruta))
    problemas: list[str] = []
    reporte: list[str] = []

    resumen = documento.tables[TABLA_RESUMEN]
    anexo = documento.tables[TABLA_ANEXO]
    filas_resumen = _filas_con_datos(resumen)
    filas_anexo = _filas_con_datos(anexo)

    total_resumen = sum(float(fila[1]) for fila in filas_resumen)
    total_anexo = sum(float(fila[6]) for fila in filas_anexo)
    total_declarado = float(resumen.rows[-1].cells[1].text.strip())
    total_anexo_declarado = float(anexo.rows[-1].cells[6].text.strip())

    reporte.append(f"Filas del resumen: {len(filas_resumen)}")
    reporte.append(f"Filas del anexo:   {len(filas_anexo)}")
    reporte.append(f"Suma 'Horas' del anexo:                {total_anexo:.1f}")
    reporte.append(f"Fila TOTAL HORAS del anexo:            {total_anexo_declarado:.1f}")
    reporte.append(f"Suma 'Dedicación (Horas)' del resumen: {total_resumen:.1f}")
    reporte.append(f"Fila TOTAL DEL PERÍODO:                {total_declarado:.1f}")

    if round(total_anexo, 2) != TOTAL_ESPERADO_HORAS:
        problemas.append(f"El anexo suma {total_anexo} y no {TOTAL_ESPERADO_HORAS}.")
    if round(total_resumen, 2) != TOTAL_ESPERADO_HORAS:
        problemas.append(f"El resumen suma {total_resumen} y no {TOTAL_ESPERADO_HORAS}.")
    if round(total_declarado, 2) != round(total_resumen, 2):
        problemas.append(
            f"TOTAL DEL PERÍODO dice {total_declarado} y el resumen suma "
            f"{total_resumen}."
        )
    if round(total_anexo_declarado, 2) != round(total_anexo, 2):
        problemas.append(
            f"TOTAL HORAS dice {total_anexo_declarado} y el anexo suma {total_anexo}."
        )

    por_persona_resumen: dict[str, float] = defaultdict(float)
    for fila in filas_resumen:
        por_persona_resumen[fila[5]] += float(fila[1])
    por_persona_anexo: dict[str, float] = defaultdict(float)
    for fila in filas_anexo:
        por_persona_anexo[fila[2]] += float(fila[6])

    reporte.append("Totales por persona (resumen | anexo | filas del anexo):")
    for persona in sorted(set(por_persona_resumen) | set(por_persona_anexo)):
        en_resumen = por_persona_resumen.get(persona, 0.0)
        en_anexo = por_persona_anexo.get(persona, 0.0)
        cuantas = sum(1 for fila in filas_anexo if fila[2] == persona)
        reporte.append(
            f"  {persona:<20} {en_resumen:>7.1f} | {en_anexo:>7.1f} | {cuantas:>3}"
        )
        if round(en_resumen, 2) != round(en_anexo, 2):
            problemas.append(
                f"{persona}: resumen {en_resumen} contra anexo {en_anexo}."
            )

    if round(por_persona_anexo.get(NOMBRE_ALEXIS, 0.0), 2) != TOTAL_ALEXIS_NETO_HORAS:
        problemas.append(
            f"{NOMBRE_ALEXIS} suma {por_persona_anexo.get(NOMBRE_ALEXIS)} y no "
            f"{TOTAL_ALEXIS_NETO_HORAS}."
        )
    dias_alexis = {fila[0] for fila in filas_anexo if fila[2] == NOMBRE_ALEXIS}
    repetidos = dias_alexis.intersection(DIAS_SIN_ALEXIS)
    for prohibido in sorted(repetidos):
        problemas.append(f"{NOMBRE_ALEXIS} tiene fila el {prohibido}.")
    reporte.append(
        f"{NOMBRE_ALEXIS} sin filas el "
        f"{' ni el '.join(DIAS_SIN_ALEXIS)}: {'no' if repetidos else 'sí'}"
    )

    filas_gabriel = [fila for fila in filas_anexo if fila[2] == NOMBRE_GABRIEL]
    if len(filas_gabriel) != FILAS_ESPERADAS_GABRIEL:
        problemas.append(
            f"{NOMBRE_GABRIEL} tiene {len(filas_gabriel)} filas y no "
            f"{FILAS_ESPERADAS_GABRIEL}."
        )
    if round(por_persona_anexo.get(NOMBRE_GABRIEL, 0.0), 2) != TOTAL_GABRIEL_HORAS:
        problemas.append(
            f"{NOMBRE_GABRIEL} suma {por_persona_anexo.get(NOMBRE_GABRIEL)} y no "
            f"{TOTAL_GABRIEL_HORAS}."
        )

    vacias = 0
    for indice, fila in enumerate(filas_anexo, start=1):
        if not fila[6] or float(fila[6]) <= 0:
            problemas.append(f"Fila {indice} del anexo con horas {fila[6]!r}.")
            vacias += 1
        if not fila[0] or not fila[2]:
            problemas.append(f"Fila {indice} del anexo sin fecha o sin persona.")
            vacias += 1
    for indice, fila in enumerate(filas_resumen, start=1):
        if not fila[1] or float(fila[1]) <= 0:
            problemas.append(f"Fila {indice} del resumen con horas {fila[1]!r}.")
            vacias += 1
    reporte.append(f"Filas con horas en cero o vacías: {vacias}")

    # Firmas y checklist tienen que haber quedado como en la plantilla.
    original = docx.Document(str(PLANTILLA))
    intactas = True
    for indice in TABLAS_INTACTAS:
        nuevo = [
            celda.text for fila in documento.tables[indice].rows for celda in fila.cells
        ]
        viejo = [
            celda.text for fila in original.tables[indice].rows for celda in fila.cells
        ]
        if nuevo != viejo:
            intactas = False
            problemas.append(f"La tabla {indice} no quedó igual que la plantilla.")
    reporte.append(
        f"Tablas {TABLAS_INTACTAS} (firmas y checklist) intactas: "
        f"{'sí' if intactas else 'NO'}"
    )

    periodo = documento.tables[TABLA_GENERALES].rows[0].cells[1].text.strip()
    if periodo != PERIODO:
        problemas.append(f"El período dice {periodo!r} y no {PERIODO!r}.")
    reporte.append(f"Período de Servicios: {periodo!r}")

    if problemas:
        raise ErrorDeRespaldo(
            "El documento generado no pasa la verificación:\n  - "
            + "\n  - ".join(problemas)
        )
    return "\n".join(reporte)


# ---------------------------------------------------------------------------


def main() -> int:
    try:
        mapeo = Mapeo.cargar(MAPEO)
        lineas_kimai = leer_kimai(mapeo)
        lineas_gabriel = leer_gabriel(mapeo)
        lineas = separar_horas_de_gabriel(lineas_kimai, lineas_gabriel)

        total = sum(linea.segundos for linea in lineas)
        if _horas(total) != TOTAL_ESPERADO_HORAS:
            raise ErrorDeRespaldo(
                f"El total del mes da {_texto_horas(total)} h y se esperaban "
                f"{TOTAL_ESPERADO_HORAS}."
            )

        filas_resumen = armar_resumen(lineas)
        filas_anexo = armar_anexo(lineas)

        # Se trabaja sobre una copia: la plantilla del dueño no se toca.
        DESTINO.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(PLANTILLA, DESTINO)
        documento = docx.Document(str(DESTINO))
        completar_documento(documento, filas_resumen, filas_anexo)
        documento.save(str(DESTINO))

        reporte = verificar_documento(DESTINO)
    except ErrorDeRespaldo as error:
        print(f"ABORTADO: {error}", file=sys.stderr)
        return 1

    print(f"Generado: {DESTINO}")
    print(reporte)
    return 0


if __name__ == "__main__":
    for flujo in (sys.stdout, sys.stderr):
        reconfigurar = getattr(flujo, "reconfigure", None)
        if reconfigurar is not None:
            try:
                reconfigurar(encoding="utf-8")
            except (ValueError, OSError):
                pass
    raise SystemExit(main())
