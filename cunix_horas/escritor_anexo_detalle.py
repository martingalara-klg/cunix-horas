"""Escritura del **Anexo II-A**, el detalle del mes en Excel.

Se parte de `templates/Anexo-II-A-Detalle-horas-KLG.xlsx`, la plantilla vacía
de C.UNIX, y se escribe una copia en `output/<mes>/`. **La plantilla nunca se
modifica**: es la fuente de la estructura, los estilos, las fórmulas y las
listas desplegables.

Qué se escribe y qué no:

- hoja `Datos`: el período, el equipo (persona, perfil, usuario de Kimai) y,
  si está configurada, la fecha del contrato. Los valores hora por proyecto
  **se leen** de ahí, no se escriben: los carga el dueño una vez.
- hoja `Detalle`: una fila por registro, columnas A a G, más la columna
  `Nota KLG` cuando alguna fila la necesita.
- hojas `Instrucciones` y `Resumen`: intactas. Son de C.UNIX.
- columnas H a M de `Detalle` (`Alertas`, `Revisión C.UNIX`, `Horas
  aprobadas`, `Horas a pagar`, `Observación C.UNIX`, `Día nuevo`): intactas.
  Las tres de cálculo ya vienen escritas como fórmula en las 600 filas del
  rango, así que alcanza con no pisarlas.

Dos cosas que openpyxl no hace solo y hay que arreglar después de guardar:

- el bloque `extLst` de la hoja `Detalle`, que son las listas desplegables de
  `Persona` y `Proyecto`. openpyxl no soporta esa extensión y la borra al
  guardar, así que se reinyecta sobre el .xlsx ya escrito.
- la verificación de integridad: el archivo se vuelve a abrir y se suma la
  columna `Horas`. Con setenta filas de seis personas en una sola hoja, una
  fila que se pierda al escribir no se ve nunca a ojo.
"""
from __future__ import annotations

import copy
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl

from cunix_horas.anexos import Periodo
from cunix_horas.filas_anexo import EPSILON, FilaAnexo

HOJA_DATOS = "Datos"
HOJA_DETALLE = "Detalle"

# La hoja `Detalle` es la cuarta del libro, y su XML interno es éste. Se usa
# para reinyectar el extLst que openpyxl descarta.
XML_DETALLE = "xl/worksheets/sheet4.xml"

# Hoja `Datos`
FILA_PERIODO = 2
FILA_CONTRATO = 3
PRIMERA_FILA_EQUIPO = 8
ULTIMA_FILA_EQUIPO = 22  # hasta donde llegan las fórmulas de `Resumen`
COL_PERSONA_DATOS, COL_PERFIL_DATOS, COL_USUARIO_DATOS = 1, 2, 3
COL_PROYECTO_DATOS, COL_CLIENTE_DATOS, COL_VALOR_HORA_DATOS = 5, 6, 7

# Hoja `Detalle`
PRIMERA_FILA = 2
ULTIMA_FILA_RANGO = 601
COL_FECHA, COL_INICIO, COL_FIN, COL_PERSONA = 1, 2, 3, 4
COL_PROYECTO, COL_DESCRIPCION, COL_HORAS = 5, 6, 7
COL_ULTIMA_CUNIX = 13  # M, la última columna de C.UNIX
COL_NOTA = 14  # N, fuera de todos los rangos de fórmulas (A..M)

ENCABEZADO_NOTA = "Nota KLG"
ANCHO_NOTA = 60


class ErrorIntegridad(Exception):
    """Lo escrito en el anexo no coincide con las horas que se leyeron."""


class ErrorCapacidad(Exception):
    """El mes no entra en la plantilla de C.UNIX."""


def leer_valores_hora(plantilla: Path) -> dict[str, tuple[str, float]]:
    """Los valores hora por proyecto que declara la hoja `Datos`.

    Devuelve proyecto -> (cliente, valor hora). Es lo que el dueño carga una
    vez en la plantilla y lo que la tabla 1 del informe usa para el importe.
    Un proyecto que no esté acá sale en el informe con sus horas y sin
    importe: inventar un valor hora sería inventar plata.
    """
    libro = openpyxl.load_workbook(plantilla, data_only=True)
    try:
        hoja = libro[HOJA_DATOS]
        valores: dict[str, tuple[str, float]] = {}
        for fila in range(PRIMERA_FILA_EQUIPO, ULTIMA_FILA_EQUIPO + 1):
            proyecto = hoja.cell(fila, COL_PROYECTO_DATOS).value
            valor = hoja.cell(fila, COL_VALOR_HORA_DATOS).value
            if not proyecto or not isinstance(valor, (int, float)):
                continue
            cliente = hoja.cell(fila, COL_CLIENTE_DATOS).value or ""
            valores[str(proyecto).strip()] = (str(cliente).strip(), float(valor))
        return valores
    finally:
        libro.close()


def _verificar_capacidad(filas, equipo) -> None:
    disponibles = ULTIMA_FILA_RANGO - PRIMERA_FILA + 1
    if len(filas) > disponibles:
        raise ErrorCapacidad(
            f"El mes tiene {len(filas)} registros y la hoja `Detalle` de la "
            f"plantilla de C.UNIX llega hasta la fila {ULTIMA_FILA_RANGO}, o "
            f"sea {disponibles} registros.\n"
            f"  No se generó nada: escribir más filas dejaría esos registros "
            f"fuera de todas las fórmulas del anexo y el total no cerraría.\n"
            f"  Pedile a C.UNIX una plantilla con más filas."
        )
    lugares = ULTIMA_FILA_EQUIPO - PRIMERA_FILA_EQUIPO + 1
    if len(equipo) > lugares:
        raise ErrorCapacidad(
            f"El mes tiene {len(equipo)} personas y la hoja `Datos` de la "
            f"plantilla de C.UNIX tiene lugar para {lugares}.\n"
            f"  No se generó nada: las personas de más quedarían fuera de la "
            f"tabla 2 del Resumen.\n"
            f"  Pedile a C.UNIX una plantilla con más filas de equipo."
        )


def _escribir_datos(hoja, periodo: Periodo, equipo, contrato_de_fecha: str) -> None:
    hoja.cell(FILA_PERIODO, 2).value = periodo.texto
    if contrato_de_fecha:
        hoja.cell(FILA_CONTRATO, 2).value = contrato_de_fecha

    for indice, (persona, perfil, usuario) in enumerate(equipo):
        fila = PRIMERA_FILA_EQUIPO + indice
        hoja.cell(fila, COL_PERSONA_DATOS).value = persona
        hoja.cell(fila, COL_PERFIL_DATOS).value = perfil
        # Vacío, no "": el anexo se compromete a que una celda vacía acá
        # signifique «esta persona todavía no tiene usuario de Kimai».
        hoja.cell(fila, COL_USUARIO_DATOS).value = usuario or None

    for fila in range(PRIMERA_FILA_EQUIPO + len(equipo), ULTIMA_FILA_EQUIPO + 1):
        for columna in (COL_PERSONA_DATOS, COL_PERFIL_DATOS, COL_USUARIO_DATOS):
            hoja.cell(fila, columna).value = None


def _escribir_detalle(hoja, filas) -> None:
    hay_notas = any(f.nota for f in filas)
    if hay_notas:
        hoja.cell(1, COL_NOTA).value = ENCABEZADO_NOTA
        hoja.cell(1, COL_NOTA)._style = copy.copy(hoja.cell(1, COL_ULTIMA_CUNIX)._style)
        hoja.column_dimensions["N"].width = ANCHO_NOTA

    for indice, f in enumerate(filas):
        fila = PRIMERA_FILA + indice
        hoja.cell(fila, COL_FECHA).value = f.fecha
        hoja.cell(fila, COL_INICIO).value = f.inicio
        hoja.cell(fila, COL_FIN).value = f.fin
        hoja.cell(fila, COL_PERSONA).value = f.persona
        hoja.cell(fila, COL_PROYECTO).value = f.proyecto
        hoja.cell(fila, COL_DESCRIPCION).value = f.descripcion or None
        hoja.cell(fila, COL_HORAS).value = f.horas
        if hay_notas:
            hoja.cell(fila, COL_NOTA).value = f.nota or None

    # Las filas sobrantes quedan vacías en A..G y N. Las fórmulas de H, K y M
    # ya vienen escritas en todo el rango 2:601 de la plantilla, incluidas las
    # filas que ahora estrenan datos: alcanza con no pisarlas.
    for fila in range(PRIMERA_FILA + len(filas), ULTIMA_FILA_RANGO + 1):
        for columna in list(range(COL_FECHA, COL_HORAS + 1)) + [COL_NOTA]:
            hoja.cell(fila, columna).value = None


def _restaurar_validaciones_x14(plantilla: Path, destino: Path) -> None:
    """Reinyecta el bloque extLst que openpyxl descarta al guardar.

    Son las listas desplegables de `Persona` y `Proyecto` de la hoja `Detalle`,
    que apuntan a la hoja `Datos`. openpyxl no soporta esa extensión y la
    elimina; sin ella, C.UNIX pierde los desplegables de su propia planilla.
    """
    with zipfile.ZipFile(plantilla) as z:
        original = z.read(XML_DETALLE).decode("utf-8")
    coincidencia = re.search(r"<extLst>.*?</extLst>", original, re.S)
    if not coincidencia:
        return
    # xr:uid sólo tiene sentido con el namespace xr, que openpyxl no declara en
    # la raíz de la hoja que genera. Sin quitarlo, el archivo no abre.
    bloque = re.sub(r'\s+xr:uid="[^"]*"', "", coincidencia.group(0))

    temporal = destino.with_name(f"~{destino.name}.extlst")
    with zipfile.ZipFile(destino) as entrada, zipfile.ZipFile(
        temporal, "w", zipfile.ZIP_DEFLATED
    ) as salida:
        for elemento in entrada.infolist():
            contenido = entrada.read(elemento.filename)
            if elemento.filename == XML_DETALLE:
                texto = contenido.decode("utf-8")
                if "<extLst>" not in texto:
                    texto = texto.replace("</worksheet>", bloque + "</worksheet>")
                contenido = texto.encode("utf-8")
            salida.writestr(elemento, contenido)
    shutil.move(str(temporal), str(destino))


def horas_escritas(ruta: Path) -> float:
    """La suma de la columna `Horas` de un anexo ya escrito."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    try:
        hoja = libro[HOJA_DETALLE]
        total = 0.0
        for fila in range(PRIMERA_FILA, ULTIMA_FILA_RANGO + 1):
            valor = hoja.cell(fila, COL_HORAS).value
            if isinstance(valor, (int, float)):
                total += float(valor)
        return total
    finally:
        libro.close()


def verificar_integridad(ruta: Path, horas_esperadas: float) -> None:
    """Relee el anexo y falla si no tiene exactamente las horas que se leyeron.

    No es un aviso: si esto no cierra, el anexo que iba a recibir C.UNIX tiene
    menos (o más) horas que los exports de Kimai, y no hay forma de saber
    cuáles a simple vista.
    """
    escritas = horas_escritas(ruta)
    if abs(escritas - horas_esperadas) <= EPSILON:
        return
    raise ErrorIntegridad(
        f"El Anexo II-A generado NO tiene las mismas horas que los exports de "
        f"Kimai: suma {escritas:.2f} h y lo leído suma {horas_esperadas:.2f} h "
        f"({escritas - horas_esperadas:+.2f} h de diferencia).\n"
        f"  Con todas las personas en una sola hoja, una fila de menos no se "
        f"ve a ojo: por eso esto frena en vez de avisar.\n"
        f"  No se generó nada. Es un error de la herramienta, no de tus "
        f"exports: guardá los archivos de input/ y avisá."
    )


def escribir(
    plantilla: Path,
    destino: Path,
    periodo: Periodo,
    filas: tuple[FilaAnexo, ...],
    equipo: tuple[tuple[str, str, str], ...],
    contrato_de_fecha: str = "",
) -> None:
    """Escribe el Anexo II-A sobre una copia de la plantilla.

    `equipo` son ternas (persona, perfil, usuario de Kimai); el usuario vacío
    es el de quien todavía no tiene cuenta.
    """
    _verificar_capacidad(filas, equipo)

    # Sin data_only: las fórmulas de la plantilla se conservan como fórmulas.
    libro = openpyxl.load_workbook(plantilla)
    try:
        _escribir_datos(libro[HOJA_DATOS], periodo, equipo, contrato_de_fecha)
        _escribir_detalle(libro[HOJA_DETALLE], filas)
        libro.save(destino)
    finally:
        libro.close()
    _restaurar_validaciones_x14(plantilla, destino)
