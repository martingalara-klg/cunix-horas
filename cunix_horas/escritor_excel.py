"""Generación del Excel con el formato que recibe el partner.

La plantilla es la fuente de estilos, no un archivo que se muta: la cantidad
de filas (proyectos por dev) y de columnas (días del mes) es variable, e
insertar filas en un .xlsx existente rompe estilos de forma silenciosa.

El estilo de cada fila no es uniforme: por ejemplo, en la fila de proyecto el
nombre va en negrita pero su total no, y los días con horas van en negrita
mientras que los días sin horas no. Por eso el estilo se modela por tipo de
fila Y por tipo de columna (etiqueta, total, día-con-valor, día-sin-valor),
tomando cada modelo de la plantilla.
"""
from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from cunix_horas.agregador import Reporte

# Nunca strftime("%b"): depende del locale de la máquina.
MESES_ABREVIADOS = {
    1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "Jun",
    7: "Jul", 8: "Aug", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dec",
}

PRIMERA_COLUMNA_DE_DIA = 3  # C
COLUMNA_ETIQUETA = 1  # A
COLUMNA_TOTAL = 2  # B

FILA_ENCABEZADO = 1
TEXTO_FILA_TOTAL = "Total"


def nombre_de_archivo(reporte: Reporte) -> str:
    """'Oct Dodera.xlsx'."""
    return f"{MESES_ABREVIADOS[reporte.mes]} {reporte.nombre_archivo}.xlsx"


def _columna_del_dia(dia: int) -> int:
    return PRIMERA_COLUMNA_DE_DIA + dia - 1


def _fila_total(hoja: Worksheet) -> int:
    """Fila cuya columna A vale 'Total'. Falla ruidosamente si no existe."""
    for fila in range(2, hoja.max_row + 1):
        if hoja.cell(row=fila, column=COLUMNA_ETIQUETA).value == TEXTO_FILA_TOTAL:
            return fila
    raise ValueError(
        "No se encontró en la plantilla ninguna fila cuya columna A "
        f"contenga el texto '{TEXTO_FILA_TOTAL}' (fila de totales)."
    )


def _fila_cliente(hoja: Worksheet, fila_total: int) -> int:
    """Fila de cliente: la que tiene las celdas de día mergeadas."""
    candidatas = sorted(
        rango.min_row
        for rango in hoja.merged_cells.ranges
        if rango.min_row == rango.max_row
        and rango.min_col == PRIMERA_COLUMNA_DE_DIA
        and 2 <= rango.min_row < fila_total
    )
    if not candidatas:
        raise ValueError(
            "No se encontró en la plantilla una fila de cliente: se esperaba "
            "una fila entre la 2 y la de 'Total' con las celdas de día "
            f"(desde la columna {get_column_letter(PRIMERA_COLUMNA_DE_DIA)}) mergeadas."
        )
    return candidatas[0]


def _filas_modelo(hoja: Worksheet) -> dict[str, int]:
    """Localiza, por estructura (no por coordenada fija), la fila modelo de
    cada tipo: encabezado, cliente, proyecto, actividad y total."""
    fila_total = _fila_total(hoja)
    fila_cliente = _fila_cliente(hoja, fila_total)
    fila_proyecto = fila_cliente + 1
    fila_actividad = fila_cliente + 2
    if fila_actividad >= fila_total:
        raise ValueError(
            "La plantilla no tiene suficientes filas entre la fila de "
            f"cliente ({fila_cliente}) y la fila '{TEXTO_FILA_TOTAL}' "
            f"({fila_total}) para identificar las filas de proyecto y actividad."
        )
    return {
        "encabezado": FILA_ENCABEZADO,
        "cliente": fila_cliente,
        "proyecto": fila_proyecto,
        "actividad": fila_actividad,
        "total": fila_total,
    }


def _celda_dia_con_y_sin_valor(hoja: Worksheet, fila: int, ultima_columna: int) -> tuple[Cell, Cell]:
    """Primera celda de día con valor y primera sin valor de la fila, buscadas
    programáticamente. Si un tipo no existe en la fila, se usa el otro."""
    con_valor: Cell | None = None
    sin_valor: Cell | None = None
    for columna in range(PRIMERA_COLUMNA_DE_DIA, ultima_columna + 1):
        celda = hoja.cell(row=fila, column=columna)
        if celda.value is not None and con_valor is None:
            con_valor = celda
        elif celda.value is None and sin_valor is None:
            sin_valor = celda
        if con_valor is not None and sin_valor is not None:
            break
    if con_valor is None and sin_valor is None:
        raise ValueError(
            f"La fila {fila} de la plantilla no tiene columnas de día "
            f"(a partir de la columna {get_column_letter(PRIMERA_COLUMNA_DE_DIA)}) "
            "de las que tomar el estilo."
        )
    return con_valor or sin_valor, sin_valor or con_valor


def _estilo_de_celda(celda: Cell) -> dict:
    return {
        "font": copy(celda.font),
        "alignment": copy(celda.alignment),
        "number_format": celda.number_format,
    }


def _estilos_de_fila(hoja: Worksheet, fila: int, ultima_columna: int) -> dict:
    """Modelos de estilo de una fila: etiqueta, total, día con valor, día sin valor."""
    dia_con, dia_sin = _celda_dia_con_y_sin_valor(hoja, fila, ultima_columna)
    return {
        "etiqueta": _estilo_de_celda(hoja.cell(row=fila, column=COLUMNA_ETIQUETA)),
        "total": _estilo_de_celda(hoja.cell(row=fila, column=COLUMNA_TOTAL)),
        "dia_con": _estilo_de_celda(dia_con),
        "dia_sin": _estilo_de_celda(dia_sin),
    }


def _estilos_de(plantilla: Path) -> dict:
    """Lee de la plantilla el estilo de cada tipo de fila, el título y los anchos."""
    libro = openpyxl.load_workbook(plantilla)
    hoja = libro.active
    ultima_columna = hoja.max_column
    filas_modelo = _filas_modelo(hoja)

    estilos: dict = {
        nombre: _estilos_de_fila(hoja, fila, ultima_columna)
        for nombre, fila in filas_modelo.items()
    }
    estilos["_titulo"] = hoja.title
    estilos["_anchos"] = {
        letra: dim.width
        for letra, dim in hoja.column_dimensions.items()
        if dim.width is not None
    }
    libro.close()
    return estilos


def _aplicar(celda: Cell, estilo_fila: dict, columna: int, tiene_valor: bool) -> None:
    if columna == COLUMNA_ETIQUETA:
        modelo = estilo_fila["etiqueta"]
    elif columna == COLUMNA_TOTAL:
        modelo = estilo_fila["total"]
    else:
        modelo = estilo_fila["dia_con"] if tiene_valor else estilo_fila["dia_sin"]
    celda.font = copy(modelo["font"])
    celda.alignment = copy(modelo["alignment"])
    celda.number_format = modelo["number_format"]


def _escribir_fila(
    hoja: Worksheet, nro_fila: int, valores: dict[int, object], estilo_fila: dict, ancho: int
) -> None:
    for columna in range(1, ancho + 1):
        celda = hoja.cell(row=nro_fila, column=columna)
        tiene_valor = columna in valores
        if tiene_valor:
            celda.value = valores[columna]
        _aplicar(celda, estilo_fila, columna, tiene_valor)


def escribir(reporte: Reporte, plantilla: Path, destino: Path) -> Path:
    """Escribe el Excel del reporte en `destino` y devuelve esa ruta."""
    estilos = _estilos_de(plantilla)
    dias = reporte.dias_del_mes
    ancho = _columna_del_dia(dias)
    ultima_letra = get_column_letter(ancho)

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = estilos["_titulo"]

    # Fila 1: nombre del dev, 'Total', y un encabezado por día como TEXTO.
    encabezado: dict[int, object] = {1: reporte.nombre_dev, 2: "Total"}
    for dia in range(1, dias + 1):
        encabezado[_columna_del_dia(dia)] = f"{reporte.mes}/{dia}/{reporte.anio}"
    _escribir_fila(hoja, 1, encabezado, estilos["encabezado"], ancho)

    nro_fila = 2
    cliente_actual: str | None = None
    proyecto_actual: tuple[str, str] | None = None

    for fila in reporte.filas:
        if fila.cliente != cliente_actual:
            total_cliente = sum(
                f.total for f in reporte.filas if f.cliente == fila.cliente
            )
            _escribir_fila(
                hoja,
                nro_fila,
                {1: fila.cliente, 2: round(total_cliente, 2)},
                estilos["cliente"],
                ancho,
            )
            hoja.merge_cells(
                f"{get_column_letter(PRIMERA_COLUMNA_DE_DIA)}{nro_fila}:"
                f"{ultima_letra}{nro_fila}"
            )
            cliente_actual = fila.cliente
            proyecto_actual = None
            nro_fila += 1

        if (fila.cliente, fila.proyecto) != proyecto_actual:
            hermanas = [
                f
                for f in reporte.filas
                if (f.cliente, f.proyecto) == (fila.cliente, fila.proyecto)
            ]
            valores: dict[int, object] = {
                1: fila.proyecto,
                2: round(sum(f.total for f in hermanas), 2),
            }
            for dia in range(1, dias + 1):
                horas = sum(f.horas_por_dia.get(dia, 0.0) for f in hermanas)
                if horas:
                    valores[_columna_del_dia(dia)] = round(horas, 2)
            _escribir_fila(hoja, nro_fila, valores, estilos["proyecto"], ancho)
            proyecto_actual = (fila.cliente, fila.proyecto)
            nro_fila += 1

        valores = {1: fila.actividad, 2: round(fila.total, 2)}
        for dia, horas in fila.horas_por_dia.items():
            valores[_columna_del_dia(dia)] = round(horas, 2)
        _escribir_fila(hoja, nro_fila, valores, estilos["actividad"], ancho)
        nro_fila += 1

    # Fila Total: con 0.0 explícito en los días sin horas, como en la plantilla.
    totales: dict[int, object] = {1: "Total", 2: round(reporte.total, 2)}
    for dia in range(1, dias + 1):
        totales[_columna_del_dia(dia)] = round(reporte.total_del_dia(dia), 2)
    _escribir_fila(hoja, nro_fila, totales, estilos["total"], ancho)

    for letra, ancho_columna in estilos["_anchos"].items():
        hoja.column_dimensions[letra].width = ancho_columna

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return destino
