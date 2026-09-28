"""Escritura del archivo de detalle plano que recibe el partner.

El contrato es `Horas KLG-Sept2025.xlsx`, el archivo que mandó el partner:
diez columnas, encabezados en negrita con relleno gris, `Date` con formato
`yyyy-mm-dd` (la hora se guarda pero no se muestra), `Duration` como
**duración real** con formato `[hh]:mm`, anchos fijos y autofiltro.

Acá no hay plantilla: a diferencia del Excel pivoteado, este archivo no tiene
estilos por tipo de fila, así que reproducir el formato es más corto y más
claro que copiarlo celda por celda de un .xlsx de muestra.

`Duration` va como `timedelta` y no como número de horas a propósito: el
partner suma esa columna, y con `[hh]:mm` una suma de duraciones da
`160:30`, mientras que un número decimal daría `160,5`. Escribirlo como
número cambiaría lo que él ve al totalizar.

La **verificación de integridad** vive acá, y relee el archivo ya escrito en
vez de confiar en las filas que se le pasaron. Ese es el punto: con 160 filas
de cinco desarrolladores en un solo archivo, una fila que se pierda al
escribir no se ve nunca a ojo, y el archivo se vería completo sin serlo.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from cunix_horas.detalle import SEGUNDOS_POR_HORA, Detalle
from cunix_horas.escritor_excel import MESES_ABREVIADOS

ENCABEZADOS = (
    "Date",
    "Duration",
    "Name",
    "User",
    "E-mail",
    "Customer",
    "Project",
    "Activity",
    "Description",
    "Project number",
)

# Anchos del archivo que mandó el partner, columna por columna de A a J.
ANCHOS = (9.22, 9.56, 11.89, 7.22, 19.67, 22.11, 21.0, 8.67, 40.67, 15.0)

RELLENO_ENCABEZADO = "FFEEEEEE"
FILA_ENCABEZADO = 1
COLUMNA_FECHA = 1  # A
COLUMNA_DURACION = 2  # B

FORMATO_FECHA = "yyyy-mm-dd"
FORMATO_DURACION = "[hh]:mm"


class ErrorIntegridad(Exception):
    """Lo escrito en el archivo no coincide con lo que se leyó de Kimai."""


def marca_de_incompleto(faltan: int) -> str:
    """' (INCOMPLETO - FALTAN 2 DESARROLLADORES - NO ENVIAR)'.

    Va en el nombre del consolidado al que le faltan desarrolladores, para que
    sea imposible mandarlo por error. Antes, cuando el entregable era un Excel
    por persona, el que faltaba se notaba solo: no estaba el archivo. Ahora
    todo va junto, el archivo se ve completo, y sin esta marca nada en la
    pantalla del dueño diría que no lo está.
    """
    if faltan == 1:
        return " (INCOMPLETO - FALTA 1 DESARROLLADOR - NO ENVIAR)"
    return f" (INCOMPLETO - FALTAN {faltan} DESARROLLADORES - NO ENVIAR)"


def nombre_de_archivo(patron: str, anio: int, mes: int) -> str:
    """'Horas KLG-Aug2026.xlsx', a partir del patrón de config/mapeo.yaml."""
    return patron.format(mes=MESES_ABREVIADOS[mes], anio=anio)


def nombre_incompleto(nombre: str, faltan: int) -> str:
    """El mismo nombre, con la marca de incompleto antes de la extensión."""
    ruta = Path(nombre)
    return f"{ruta.stem}{marca_de_incompleto(faltan)}{ruta.suffix}"


def escribir_detalle(detalle: Detalle, destino: Path) -> None:
    """Escribe el archivo del partner con las filas ya ordenadas."""
    libro = openpyxl.Workbook()
    hoja = libro.active

    negrita = Font(bold=True)
    relleno = PatternFill(
        fill_type="solid",
        start_color=RELLENO_ENCABEZADO,
        end_color=RELLENO_ENCABEZADO,
    )
    for indice, texto in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=FILA_ENCABEZADO, column=indice, value=texto)
        celda.font = negrita
        celda.fill = relleno

    for indice, ancho in enumerate(ANCHOS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    for numero, fila in enumerate(detalle.filas, start=FILA_ENCABEZADO + 1):
        for indice, valor in enumerate(fila.valores, start=1):
            celda = hoja.cell(row=numero, column=indice, value=valor)
            if indice == COLUMNA_FECHA:
                celda.number_format = FORMATO_FECHA
            elif indice == COLUMNA_DURACION:
                celda.number_format = FORMATO_DURACION

    ultima = FILA_ENCABEZADO + len(detalle.filas)
    hoja.auto_filter.ref = (
        f"A{FILA_ENCABEZADO}:{get_column_letter(len(ENCABEZADOS))}{ultima}"
    )

    libro.save(destino)


def segundos_escritos(ruta: Path) -> int:
    """La suma de la columna Duration del archivo ya escrito, en segundos."""
    libro = openpyxl.load_workbook(ruta)
    try:
        hoja = libro.active
        total = 0
        for numero in range(FILA_ENCABEZADO + 1, hoja.max_row + 1):
            valor = hoja.cell(row=numero, column=COLUMNA_DURACION).value
            if valor is None:
                continue
            total += int(valor.total_seconds())
        return total
    finally:
        libro.close()


def verificar_integridad(ruta: Path, segundos_esperados: int) -> None:
    """Relee el archivo y falla si no tiene exactamente las horas leídas.

    No es un aviso: si esto no cierra, el archivo que iba a recibir el partner
    tiene menos (o más) horas que los exports de Kimai, y no hay forma de
    saber cuáles a simple vista. Se falla ruidoso y no se reemplaza nada.
    """
    escritos = segundos_escritos(ruta)
    if escritos == segundos_esperados:
        return
    diferencia = (escritos - segundos_esperados) / SEGUNDOS_POR_HORA
    raise ErrorIntegridad(
        f"El archivo generado NO tiene las mismas horas que los exports de "
        f"Kimai: suma {escritos / SEGUNDOS_POR_HORA:.2f} h y los registros "
        f"leídos suman {segundos_esperados / SEGUNDOS_POR_HORA:.2f} h "
        f"({diferencia:+.2f} h de diferencia).\n"
        f"  Con todos los desarrolladores en un solo archivo, una fila de "
        f"menos no se ve a ojo: por eso esto frena en vez de avisar.\n"
        f"  No se generó nada. Es un error de la herramienta, no de tus "
        f"exports: guardá los archivos de input/ y avisá."
    )
