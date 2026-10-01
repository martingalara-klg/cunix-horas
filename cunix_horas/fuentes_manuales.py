"""Las planillas de `input/AAAA-MM/manual/`: lo que no sale de Kimai.

Son dos casos, los dos declarados en `config/mapeo.yaml` bajo
`fuentes_manuales:`, y los dos opcionales: un mes sin fuentes declaradas corre
igual.

**a) Horas de alguien que no está en Kimai.** La persona trabajó, pero no
tiene usuario, así que sus horas quedaron cargadas en la cuenta de otro. La
planilla dice cuántas horas hizo cada día; esas horas se le restan día por día
a la otra persona y pasan a filas propias. Si la resta no cierra —un día en
que la otra persona no tiene horas, o no le alcanzan— **no se genera nada**:
un reparto que no cuadra se factura mal y nadie lo ve en una hoja de setenta
filas.

**b) Descripciones que alguien entrega aparte.** Desde septiembre de 2026
todos cargan la descripción en Kimai, así que esto es la excepción, no el
camino normal. La planilla trae las mismas columnas que la hoja `Detalle` y
sólo aporta la descripción: el día y las horas se verifican contra lo que ya
está, y si algo no coincide tampoco se genera nada.

**Las planillas se leen de forma tolerante.** No se asume ni el nombre de la
hoja, ni el número de fila del encabezado, ni el formato de la fecha: el
encabezado se busca por el nombre de las columnas y la fecha se acepta como
texto en español sin año (`Dom 02 Ago`), como `dd/mm/aaaa`, como `aaaa-mm-dd`
o como fecha de Excel. Si aun así no se entiende, el error dice qué se
esperaba encontrar y qué se encontró en su lugar.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, replace
from pathlib import Path

import openpyxl

from cunix_horas.anexos import (
    DescripcionesAparte,
    FuentesManuales,
    HorasSinKimai,
    Periodo,
)
from cunix_horas.filas_anexo import EPSILON, FilaAnexo, ordenar, restar

# Las declara `anexos.py` para que `mapeo.py` pueda leerlas del YAML sin
# depender de este módulo (que a su vez depende del mapeo). Se re-exportan acá
# porque es acá donde se usan.
__all__ = [
    "CARPETA_MANUAL",
    "DescripcionesAparte",
    "ErrorFuenteManual",
    "FuentesManuales",
    "HorasSinKimai",
    "RenglonManual",
    "aplicar",
    "leer_planilla_de_descripciones",
    "leer_planilla_de_horas",
    "nota_por_defecto",
]

# Hasta qué fila se busca el encabezado. Estas planillas suelen traer un título
# y un par de líneas de contexto arriba.
FILAS_DE_BUSQUEDA = 20

CARPETA_MANUAL = "manual"

_SOLO_DIA = re.compile(r"\b(\d{1,2})\b")
_FECHA_CON_BARRAS = re.compile(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\s*$")
_FECHA_ISO = re.compile(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})\s*$")
_HORA = re.compile(r"^\s*(\d{1,2}):([0-5]\d)(?::[0-5]\d)?\s*$")

# Nombres con los que se reconoce cada columna. Se compara sin acentos, sin
# mayúsculas y por prefijo, para que «Descripción (ticket – qué se hizo)»
# coincida con «descripcion».
ALIAS = {
    "fecha": ("fecha", "dia"),
    "horas": ("horas", "hs"),
    "descripcion": ("descripcion", "detalle", "trabajo"),
    "inicio": ("inicio", "desde"),
    "fin": ("fin", "hasta"),
}

_ACENTOS = str.maketrans("áéíóúüñ", "aeiouun")


class ErrorFuenteManual(Exception):
    """Una planilla manual no se pudo leer, o lo que trae no cierra.

    Frena la corrida entera a propósito: estas planillas mueven horas de una
    persona a otra, y generar los anexos con un reparto que no cuadra es peor
    que no generarlos.
    """


def nota_por_defecto(conf: HorasSinKimai) -> str:
    """La `Nota KLG` de las filas que mueve una fuente manual."""
    return (
        f"Horas de {conf.persona}, que no tiene usuario de Kimai: quedaron "
        f"cargadas en la cuenta de {conf.restar_a} y KLG las separó día por "
        f"día a partir de la planilla {conf.planilla}. Las horas de cada uno "
        f"están respaldadas; el corte dentro del día es una reconstrucción."
    )


def _normalizar(texto) -> str:
    return str(texto or "").strip().lower().translate(_ACENTOS)


def _columna_de(celdas: dict[int, str], clave: str) -> int | None:
    for indice in sorted(celdas):
        for alias in ALIAS[clave]:
            if celdas[indice].startswith(alias):
                return indice
    return None


def _ubicar_encabezado(ruta: Path, requeridas: tuple[str, ...], opcionales: tuple[str, ...]):
    """Busca, en todas las hojas, la fila que tiene las columnas pedidas.

    Devuelve (hoja, número de fila del encabezado, {clave: columna}).
    """
    libro = openpyxl.load_workbook(ruta, data_only=True)
    vistas: list[str] = []
    for hoja in libro.worksheets:
        tope = min(hoja.max_row or 0, FILAS_DE_BUSQUEDA)
        for fila in range(1, tope + 1):
            celdas = {
                c: _normalizar(hoja.cell(fila, c).value)
                for c in range(1, (hoja.max_column or 0) + 1)
            }
            celdas = {c: t for c, t in celdas.items() if t}
            if not celdas:
                continue
            vistas.append(
                f"hoja «{hoja.title}», fila {fila}: {list(celdas.values())}"
            )
            columnas = {
                clave: _columna_de(celdas, clave)
                for clave in requeridas + opcionales
            }
            if all(columnas[clave] is not None for clave in requeridas):
                return hoja, fila, columnas

    esperadas = ", ".join(f"«{ALIAS[c][0]}»" for c in requeridas)
    raise ErrorFuenteManual(
        f"No se pudo interpretar la planilla manual {ruta.name}.\n"
        f"  Se esperaba una fila de encabezados con, al menos, las columnas "
        f"{esperadas} (no importan el orden, los acentos, las mayúsculas ni en "
        f"qué hoja estén).\n"
        f"  Se buscó en las primeras {FILAS_DE_BUSQUEDA} filas de cada hoja y "
        f"se encontró esto:\n"
        + "\n".join(f"    {v}" for v in (vistas[:10] or ["(ninguna fila con texto)"]))
        + f"\n  Corregí los encabezados de {ruta.name} y volvé a correr."
    )


def _error_de_otro_mes(texto: str, periodo: Periodo, donde: str) -> ErrorFuenteManual:
    return ErrorFuenteManual(
        f"{donde}: la fecha «{texto}» no cae en {periodo.texto}.\n"
        f"  Esa planilla manual es de otro mes: conseguí la del mes que estás "
        f"generando."
    )


def _dia_de(valor, periodo: Periodo, donde: str) -> int | None:
    """El día del mes que indica una celda de fecha, o None si no es una fecha.

    Acepta la fecha de Excel, 'dd/mm/aaaa', 'aaaa-mm-dd' y el texto en español
    sin año que usan las planillas hechas a mano ('Dom 02 Ago').
    """
    if isinstance(valor, dt.datetime):
        valor = valor.date()
    if isinstance(valor, dt.date):
        if not periodo.contiene(valor):
            raise _error_de_otro_mes(f"{valor:%d/%m/%Y}", periodo, donde)
        return valor.day

    texto = str(valor or "").strip()
    if not texto:
        return None

    coincidencia = _FECHA_CON_BARRAS.match(texto)
    if coincidencia:
        dia, mes, anio = (int(g) for g in coincidencia.groups())
        anio += 2000 if anio < 100 else 0
        if (anio, mes) != (periodo.anio, periodo.mes):
            raise _error_de_otro_mes(texto, periodo, donde)
        return dia

    coincidencia = _FECHA_ISO.match(texto)
    if coincidencia:
        anio, mes, dia = (int(g) for g in coincidencia.groups())
        if (anio, mes) != (periodo.anio, periodo.mes):
            raise _error_de_otro_mes(texto, periodo, donde)
        return dia

    # Último recurso: un texto en español sin año, del estilo 'Dom 02 Ago'. El
    # día es lo único que hace falta; el mes y el año son los del período.
    coincidencia = _SOLO_DIA.search(texto)
    if not coincidencia:
        return None
    dia = int(coincidencia.group(1))
    return dia if 1 <= dia <= periodo.dias else None


def _horas_de(valor, donde: str) -> float | None:
    if isinstance(valor, bool):
        raise ErrorFuenteManual(f"{donde}: «{valor}» no es una cantidad de horas.")
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor or "").strip().replace(",", ".")
    if not texto:
        return None
    try:
        return float(texto)
    except ValueError:
        raise ErrorFuenteManual(
            f"{donde}: «{valor}» no es una cantidad de horas."
        ) from None


def _hora_de(valor) -> dt.time | None:
    if isinstance(valor, dt.datetime):
        return valor.time()
    if isinstance(valor, dt.time):
        return valor
    coincidencia = _HORA.match(str(valor or ""))
    if not coincidencia:
        return None
    return dt.time(int(coincidencia.group(1)), int(coincidencia.group(2)))


@dataclass(frozen=True)
class RenglonManual:
    """Un renglón ya interpretado de una planilla manual."""

    dia: int
    horas: float | None
    descripcion: str
    inicio: dt.time | None = None
    fin: dt.time | None = None


def _leer_renglones(
    ruta: Path,
    periodo: Periodo,
    requeridas: tuple[str, ...],
    opcionales: tuple[str, ...],
) -> tuple[RenglonManual, ...]:
    """Los renglones de datos de una planilla manual, salteando lo que no lo es.

    Un renglón entra si su celda de fecha se entiende como un día del mes. Las
    filas de subtotal y de total quedan afuera solas: no tienen fecha.
    """
    hoja, fila_encabezado, columnas = _ubicar_encabezado(ruta, requeridas, opcionales)
    renglones: list[RenglonManual] = []
    vistos: set[int] = set()

    for fila in range(fila_encabezado + 1, (hoja.max_row or 0) + 1):
        donde = f"{ruta.name}, hoja «{hoja.title}», fila {fila}"
        dia = _dia_de(hoja.cell(fila, columnas["fecha"]).value, periodo, donde)
        if dia is None:
            continue
        if dia in vistos:
            raise ErrorFuenteManual(
                f"{donde}: el día {dia} aparece más de una vez en {ruta.name}.\n"
                f"  Esta planilla tiene que traer un solo renglón por día."
            )
        vistos.add(dia)

        renglones.append(
            RenglonManual(
                dia=dia,
                horas=(
                    _horas_de(hoja.cell(fila, columnas["horas"]).value, donde)
                    if columnas.get("horas") is not None
                    else None
                ),
                descripcion=(
                    str(hoja.cell(fila, columnas["descripcion"]).value or "").strip()
                    if columnas.get("descripcion") is not None
                    else ""
                ),
                inicio=(
                    _hora_de(hoja.cell(fila, columnas["inicio"]).value)
                    if columnas.get("inicio") is not None
                    else None
                ),
                fin=(
                    _hora_de(hoja.cell(fila, columnas["fin"]).value)
                    if columnas.get("fin") is not None
                    else None
                ),
            )
        )

    if not renglones:
        raise ErrorFuenteManual(
            f"La planilla manual {ruta.name} no tiene ningún renglón de datos.\n"
            f"  El encabezado se encontró en la fila {fila_encabezado} de la "
            f"hoja «{hoja.title}», pero debajo no hay ninguna fila cuya celda "
            f"de fecha se entienda como un día de {periodo.texto}.\n"
            f"  Se aceptan la fecha de Excel, «dd/mm/aaaa», «aaaa-mm-dd» y el "
            f"texto en español sin año («Dom 02 Ago»)."
        )
    return tuple(renglones)


def leer_planilla_de_horas(ruta: Path, periodo: Periodo) -> tuple[RenglonManual, ...]:
    """Los días y las horas de quien no tiene usuario de Kimai."""
    renglones = _leer_renglones(ruta, periodo, ("fecha", "horas"), ("descripcion",))
    sin_horas = [r.dia for r in renglones if r.horas is None or r.horas <= 0]
    if sin_horas:
        raise ErrorFuenteManual(
            f"En {ruta.name} hay días sin horas o con horas en cero: "
            f"{', '.join(str(d) for d in sin_horas)}.\n"
            f"  Cada renglón tiene que decir cuántas horas se trabajaron ese día."
        )
    return renglones


def leer_planilla_de_descripciones(
    ruta: Path, periodo: Periodo
) -> tuple[RenglonManual, ...]:
    """Las descripciones que alguien entrega fuera de Kimai."""
    return _leer_renglones(
        ruta, periodo, ("fecha", "descripcion"), ("horas", "inicio", "fin")
    )


def _ruta_de(carpeta: Path, planilla: str, que_es: str, persona: str) -> Path:
    ruta = carpeta / planilla
    if ruta.is_file():
        return ruta
    raise ErrorFuenteManual(
        f"Falta la planilla manual {planilla}, de {persona}.\n"
        f"  Se esperaba encontrarla en {carpeta}\n"
        f"  config/mapeo.yaml la declara bajo fuentes_manuales: ({que_es}), "
        f"así que sin ella el mes no se puede armar: NO SE GENERÓ NADA.\n"
        f"  Conseguí la planilla y dejala ahí. Si esa persona ya carga sus "
        f"horas en Kimai, sacá su entrada de config/mapeo.yaml."
    )


def _separar_horas(
    filas: tuple[FilaAnexo, ...],
    usuarios: dict[str, str],
    conf: HorasSinKimai,
    carpeta: Path,
    periodo: Periodo,
) -> tuple[FilaAnexo, ...]:
    """Le resta a `conf.restar_a` las horas de `conf.persona`, día por día.

    El corte dentro del día es por horario: la primera parte del bloque queda
    para quien lo tenía cargado y la última pasa a la persona de la planilla.
    Es una reconstrucción —Kimai registró un solo bloque— y queda dicho en la
    columna `Nota KLG` de las filas afectadas.
    """
    ruta = _ruta_de(carpeta, conf.planilla, "horas_sin_kimai", conf.persona)
    renglones = leer_planilla_de_horas(ruta, periodo)

    if any(f.persona == conf.persona for f in filas):
        raise ErrorFuenteManual(
            f"{conf.persona} ya tiene horas propias en los exports de Kimai de "
            f"{periodo.texto}, y config/mapeo.yaml sigue declarando que sus "
            f"horas vienen de {conf.planilla}.\n"
            f"  Si ya tiene usuario de Kimai, sacá su entrada de "
            f"fuentes_manuales: en config/mapeo.yaml. Si no, sus horas se "
            f"contarían dos veces: NO SE GENERÓ NADA."
        )

    por_dia: dict[int, list[FilaAnexo]] = {}
    for fila in filas:
        if fila.persona == conf.restar_a:
            por_dia.setdefault(fila.fecha.day, []).append(fila)

    problemas = []
    for renglon in renglones:
        disponibles = por_dia.get(renglon.dia, [])
        if not disponibles:
            problemas.append(
                f"El {renglon.dia}/{periodo.mes:02d} {conf.persona} tiene "
                f"{renglon.horas:g} h en {conf.planilla} y {conf.restar_a} no "
                f"tiene ninguna hora cargada ese día."
            )
            continue
        total = sum(f.horas for f in disponibles)
        if total + EPSILON < renglon.horas:
            problemas.append(
                f"El {renglon.dia}/{periodo.mes:02d} {conf.persona} tiene "
                f"{renglon.horas:g} h en {conf.planilla} y {conf.restar_a} sólo "
                f"tiene {total:g} h cargadas: el día quedaría en negativo."
            )
    if problemas:
        raise ErrorFuenteManual(
            f"La resta de las horas de {conf.persona} a {conf.restar_a} no "
            f"cierra. NO SE GENERÓ NADA:\n"
            + "\n".join(f"  - {p}" for p in problemas)
            + f"\n  Revisá {conf.planilla} contra el export de Kimai de "
            f"{conf.restar_a} y volvé a correr."
        )

    nota = conf.nota or nota_por_defecto(conf)
    descripciones = {r.dia: r.descripcion for r in renglones}
    a_restar = {r.dia: r.horas for r in renglones}

    resultado: list[FilaAnexo] = [f for f in filas if f.persona != conf.restar_a]
    for dia, disponibles in por_dia.items():
        pendiente = a_restar.get(dia, 0.0)
        if pendiente <= EPSILON:
            resultado.extend(disponibles)
            continue
        for fila in reversed(disponibles):
            if pendiente <= EPSILON:
                resultado.append(fila)
                continue
            toma = min(fila.horas, pendiente)
            pendiente -= toma
            entero = toma + EPSILON >= fila.horas
            corte = restar(fila.fin, toma) if fila.fin is not None else None
            resultado.append(
                FilaAnexo(
                    fecha=fila.fecha,
                    inicio=fila.inicio if entero else corte,
                    fin=fila.fin,
                    persona=conf.persona,
                    proyecto=conf.proyecto or fila.proyecto,
                    descripcion=descripciones.get(dia, ""),
                    horas=toma,
                    nota=nota,
                )
            )
            if not entero:
                resultado.append(
                    replace(fila, horas=fila.horas - toma, fin=corte, nota=nota)
                )

    usuarios.setdefault(conf.persona, "")
    return ordenar(resultado)


def _cargar_descripciones(
    filas: tuple[FilaAnexo, ...],
    conf: DescripcionesAparte,
    carpeta: Path,
    periodo: Periodo,
) -> tuple[FilaAnexo, ...]:
    """Carga en las filas de una persona las descripciones que entregó aparte."""
    ruta = _ruta_de(carpeta, conf.planilla, "descripciones", conf.persona)
    renglones = leer_planilla_de_descripciones(ruta, periodo)

    por_dia: dict[int, list[FilaAnexo]] = {}
    for fila in filas:
        if fila.persona == conf.persona:
            por_dia.setdefault(fila.fecha.day, []).append(fila)

    problemas = []
    for renglon in renglones:
        presentes = por_dia.get(renglon.dia, [])
        if not presentes:
            problemas.append(
                f"La planilla tiene el día {renglon.dia}/{periodo.mes:02d} y "
                f"{conf.persona} no tiene horas cargadas ese día."
            )
            continue
        if renglon.horas is not None:
            total = sum(f.horas for f in presentes)
            if abs(total - renglon.horas) > EPSILON:
                problemas.append(
                    f"El {renglon.dia}/{periodo.mes:02d} la planilla dice "
                    f"{renglon.horas:g} h y el anexo tiene {total:g} h."
                )
    if problemas:
        raise ErrorFuenteManual(
            f"La planilla de descripciones de {conf.persona} "
            f"({conf.planilla}) no coincide con sus horas. NO SE GENERÓ NADA:\n"
            + "\n".join(f"  - {p}" for p in problemas)
            + "\n  Esa planilla sólo puede aportar descripciones: el día y las "
            "horas tienen que ser los mismos que ya tiene el anexo."
        )

    textos = {r.dia: r.descripcion for r in renglones if r.descripcion}
    horarios = {r.dia: (r.inicio, r.fin) for r in renglones}

    resultado = []
    for fila in filas:
        if fila.persona != conf.persona or fila.fecha.day not in textos:
            resultado.append(fila)
            continue
        inicio, fin = horarios.get(fila.fecha.day, (None, None))
        # El horario sólo se toma cuando el export no lo trae y la planilla da
        # uno solo para ese día. Si el anexo ya tiene horario, ése manda.
        unica = len(por_dia[fila.fecha.day]) == 1
        resultado.append(
            replace(
                fila,
                descripcion=textos[fila.fecha.day],
                inicio=inicio if (unica and fila.inicio is None and inicio) else fila.inicio,
                fin=fin if (unica and fila.fin is None and fin) else fila.fin,
            )
        )
    return ordenar(resultado)


def aplicar(
    filas: tuple[FilaAnexo, ...],
    usuarios: tuple[tuple[str, str], ...],
    fuentes: FuentesManuales,
    carpeta_entrada: Path,
    periodo: Periodo,
) -> tuple[tuple[FilaAnexo, ...], tuple[tuple[str, str], ...]]:
    """Aplica todas las fuentes manuales declaradas. No muta lo que recibe.

    Lanza `ErrorFuenteManual` ante cualquier planilla que falte, que no se
    entienda o que no cierre: en ese caso la corrida no genera ningún anexo.
    """
    carpeta = carpeta_entrada / CARPETA_MANUAL
    nuevos = dict(usuarios)

    for conf in fuentes.horas_sin_kimai:
        filas = _separar_horas(filas, nuevos, conf, carpeta, periodo)
    for conf in fuentes.descripciones:
        filas = _cargar_descripciones(filas, conf, carpeta, periodo)

    return filas, tuple(sorted(nuevos.items()))
